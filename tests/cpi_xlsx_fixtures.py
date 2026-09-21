"""Tiny MoSPI-shaped workbooks for ingest tests. Not producer bytes."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from prism.ingest.xlsx_cpi_period import (
    BACK_COLUMNS,
    DIVISION_COLUMNS,
    GENERAL_COLUMNS,
    GROUP_COLUMNS,
)


def _xlsx(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def monthly_workbook_bytes() -> bytes:
    workbook = Workbook()
    item = workbook.active
    assert item is not None
    item.title = "Item"
    item.append(["state_name", "index"])
    item.append(["should-not-appear", 999])

    subclass = workbook.create_sheet("Subclass")
    subclass.append(["unused"])
    class_sheet = workbook.create_sheet("Class")
    class_sheet.append(["unused"])

    group = workbook.create_sheet("Group")
    group.append(list(GROUP_COLUMNS))
    group.append(
        [
            "2024",
            "00",
            "All India",
            "Rural",
            "Food",
            "01.1",
            2026,
            "July",
            108.35,
            4.54,
            "F",
        ]
    )
    group.append(
        [
            "2024",
            "00",
            "All India",
            "Urban",
            "Food",
            "01.1",
            2026,
            "July",
            107.0,
            4.0,
            "F",
        ]
    )
    group.append(
        [
            "2024",
            "00",
            "All India",
            "Rural",
            "Beverages",
            "01.2",
            2026,
            "July",
            107.06,
            1.45,
            "F",
        ]
    )
    group.append(
        [
            "2024",
            "04",
            "Chandigarh",
            "Urban",
            "Food",
            "01.1",
            2026,
            "July",
            111.11,
            5.39,
            "F",
        ]
    )
    group.append(
        [
            "2024",
            "04",
            "Chandigarh",
            "Combined",
            "Food",
            "01.1",
            2026,
            "July",
            111.11,
            5.39,
            "F",
        ]
    )

    division = workbook.create_sheet("Division")
    division.append(list(DIVISION_COLUMNS))
    division.append(
        [
            "2024",
            "00",
            "All India",
            "Rural",
            "Food and beverages",
            "01",
            2026,
            "July",
            108.27,
            4.33,
            "F",
        ]
    )
    division.append(
        [
            "2024",
            "00",
            "All India",
            "Rural",
            "Clothing and footwear",
            "03",
            2026,
            "July",
            108.31,
            2.65,
            "F",
        ]
    )

    general = workbook.create_sheet("General")
    general.append(list(GENERAL_COLUMNS))
    general.append(
        ["2024", "00", "All India", "Rural", 2026, "July", 107.94, 3.74, "F"]
    )
    general.append(
        ["2024", "00", "All India", "Urban", 2026, "July", 107.22, 4.23, "F"]
    )
    general.append(
        ["2024", "00", "All India", "Combined", 2026, "July", 107.6, 3.96, "F"]
    )
    general.append(
        ["2024", "04", "Chandigarh", "Urban", 2026, "July", 108.04, 4.3, "F"]
    )
    general.append(
        ["2024", "04", "Chandigarh", "Combined", 2026, "July", 108.04, 4.3, "F"]
    )
    general.append(
        ["2024", "01", "Jammu And Kashmir", "Rural", 2026, "July", 107.94, "-", "F"]
    )

    return _xlsx(workbook)


def annex_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Annexure-III"
    sheet.append(["State/UT wise general index"])
    sheet.append(["S.No.", "State", "Rural"])
    sheet.append(["6", "Chandigarh*", None])
    sheet.append(["*: No rural market in Chandigarh"])
    return _xlsx(workbook)


def back_series_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Sheet1"
    sheet.append(list(BACK_COLUMNS))
    sheet.append(
        ["2024", "00", "All India", "Rural", 2013, "January", "General", 54.9, None]
    )
    sheet.append(
        ["2024", "00", "All India", "Urban", 2013, "January", "General", 55.3, None]
    )
    sheet.append(
        ["2024", "00", "All India", "Combined", 2013, "January", "General", 55.1, None]
    )
    sheet.append(
        ["2024", "00", "All India", "Rural", 2014, "January", "General", 60.0, 9.29]
    )
    return _xlsx(workbook)
