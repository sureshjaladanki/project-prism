"""Tiny Budget/CGA-shaped artifacts for C3 ingest tests. Not producer bytes."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from prism.ingest.parse_c3 import DEFICIT_SHEET, RECEIPT_SHEET, STAT1_SHEET


def _xlsx(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def receipt_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = RECEIPT_SHEET
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append(
        [
            None,
            "Tax Revenue",
            None,
            None,
            None,
            "Major Head",
            "Actuals\n2024-2025",
            "Budget\n2025-2026",
            "Revised \n2025-2026",
            "Budget\n2026-2027",
        ]
    )
    sheet.append(
        [None, "Corporation Tax", None, None, None, "0020", 911000, 1000000, 980000, 1100000]
    )
    sheet.append([None, "CGST", None, None, None, "0005", 800000, "", 810000, 900000])
    return _xlsx(workbook)


def expenditure_stat1_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = STAT1_SHEET
    for _ in range(4):
        sheet.append([])
    year = [None] * 17
    year[4] = "Actuals 2024-2025"
    year[7] = "Budget Estimates 2025-2026"
    year[10] = "Revised Estimates 2025-2026"
    year[13] = "Budget Estimates 2026-2027"
    sheet.append(year)
    split = [None] * 17
    for start in (4, 7, 10):
        split[start] = "Revenue"
        split[start + 1] = "Capital"
        split[start + 2] = "Total"
    split[13] = "Revenue"
    split[14] = "Capital"
    split[16] = "Total"
    sheet.append(split)
    data = [None] * 17
    data[3] = "Central Expenditure (2+3+4)"
    values = (
        1,
        2,
        3,
        4,
        5,
        6,
        7,
        8,
        9,
        10,
        11,
        13,
    )
    indexes = (4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 16)
    for index, value in zip(indexes, values, strict=True):
        data[index] = value
    sheet.append(data)
    return _xlsx(workbook)


def deficit_statistics_bytes() -> bytes:
    workbook = Workbook()
    decoy = workbook.active
    assert decoy is not None
    decoy.title = "BAAG"
    decoy.append(["should-not-appear", 999])
    sheet = workbook.create_sheet(DEFICIT_SHEET)
    sheet.append([])
    sheet.append(["Deficit Statistics"])
    sheet.append([])
    years = [None] * 7
    years[3] = "2024-2025"
    years[4] = "2025-2026"
    years[5] = "2025-2026"
    years[6] = "2026-2027"
    sheet.append(years)
    status = [None] * 7
    status[3] = "वास्तविक\n\nActuals"
    status[4] = "बजट\nअनुमान\nBudget \nEstimates"
    status[5] = "संशोधित\nअनुमान\nRevised \nEstimates"
    status[6] = "बजट\nअनुमान\nBudget \nEstimates"
    sheet.append(status)
    row = [None] * 7
    row[1] = "राजकोषीय घाटा"
    row[2] = " 1. Fiscal Deficit"
    row[3] = 1574431
    row[4] = 1568936
    row[5] = 1558492
    row[6] = 1695768
    sheet.append(row)
    return _xlsx(workbook)


def cga_monthly_html_bytes() -> bytes:
    return """<!DOCTYPE html>
<html>
<body>
<table>
<tr>
<td></td><td></td><td></td>
<td>Budget Estimates 2026-2027</td>
<td>Actuals @ upto July 2026</td>
<td>% of Actuals to Budget Estimates</td>
<td>% of Actuals to Budget Estimates</td>
</tr>
<tr>
<td></td><td></td><td></td>
<td>Rs.</td><td>Rs.</td><td>Current</td><td>COPPY</td>
</tr>
<tr>
<td>1</td><td>Revenue Receipts</td><td></td>
<td>3533150</td><td>1267573</td><td>35.9 %</td><td>(31.1%)</td>
</tr>
</table>
</body>
</html>
""".encode("cp1252")


STUB_PDF = b"%PDF-1.7\n1 0 obj<<>>endobj\n%%EOF"
