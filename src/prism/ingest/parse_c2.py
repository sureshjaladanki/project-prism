"""Tidy C2 producer tables. Stop when the artifact cannot be read as cells."""

from __future__ import annotations

import csv
import re
from collections.abc import Callable
from io import BytesIO, StringIO
from typing import Any

import pdfplumber
from openpyxl import load_workbook
from pdfplumber.utils.exceptions import PdfminerException

from prism.ingest.parse import ParsedTable
from prism.ingest.parse_c2_layout import XlsSheet, read_xls_sheets
from prism.ingest.parse_pdf_layout import (
    AssignedTable,
    PdfColumn,
    PdfWord,
    assign_pdf_table,
    cluster_rows,
    group_centers,
    is_amount_token,
    names_from_groups,
    page_words,
    words_before,
    words_through,
)
from prism.ingest.retrieve import is_xls_ole
from prism.schema import YesNo

PARSER_NAME = "orgi-c2"
PARSER_VERSION = "1.2.0"
PARSER = f"{PARSER_NAME}-{PARSER_VERSION}"

PCA_COLUMNS = (
    "State",
    "District",
    "Level",
    "Name",
    "TRU",
    "No_HH",
    "TOT_P",
    "TOT_M",
    "TOT_F",
)
PCA_LEVELS = frozenset({"India", "STATE"})
A02_COLUMNS = (
    "state_code",
    "district_code",
    "name",
    "census_year",
    "persons",
    "variation_absolute",
    "variation_percentage",
    "males",
    "females",
)
A02_VALUE_COLUMNS = A02_COLUMNS[4:]
A02_SHEET = "A-2"
PDF_Y_TOL = 5.0
GROUP_WIDTH = 3
BULLETIN_VALUE_COLUMNS = (
    "birth_total",
    "birth_rural",
    "birth_urban",
    "death_total",
    "death_rural",
    "death_urban",
    "ngr_total",
    "ngr_rural",
    "ngr_urban",
    "imr_total",
    "imr_rural",
    "imr_urban",
)
BULLETIN_COLUMNS = ("line_label",) + BULLETIN_VALUE_COLUMNS
BULLETIN_PDF_COLUMNS = tuple(
    PdfColumn(name, token)
    for name, token in zip(
        BULLETIN_VALUE_COLUMNS, ("Total", "Rural", "Urban") * 4, strict=True
    )
)
TABLE3_COLUMNS = ("geography", "line_label", "total", "rural", "urban")
TABLE3_VALUE_COLUMNS = ("total", "rural", "urban")
TABLE8_COLUMNS = ("geography", "year", "persons", "male", "females")
TABLE8_VALUE_COLUMNS = ("persons", "male", "females")
TABLE8_SKIP_WORDS = frozenset(
    {
        "Year",
        "TABLE",
        "–",
        "-",
        "8",
        "(Contd.)",
        "Projected",
        "Total",
        "Population",
        "by",
        "Sex",
        "as",
        "on",
        "1st",
        "March",
        ":",
        "India,",
        "States",
        "and",
        "Union",
        "Territories*",
        "('000)",
    }
)
_YEAR = re.compile(r"(?:19|20)\d{2}")
_WITHHELD = frozenset(
    {"-", "–", "—", "---", "n.a", "n.a.", "na", "n/a", "not available"}
)


def _empty_fail(flags: str) -> ParsedTable:
    return ParsedTable(
        csv_text="",
        row_count=0,
        nulls="not parsed",
        flags=flags,
        lineage_ok=YesNo.no,
    )


def _to_csv(fieldnames: tuple[str, ...], rows: list[dict[str, str]]) -> str:
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def _cell(value: str) -> str:
    text = value.strip()
    if text.lower() in _WITHHELD:
        return ""
    return text


def _ok_rows(
    fieldnames: tuple[str, ...],
    rows: list[dict[str, str]],
    value_columns: tuple[str, ...],
    flags: str,
) -> ParsedTable:
    nulls = sum(1 for row in rows for name in value_columns if row[name] == "")
    return ParsedTable(
        csv_text=_to_csv(fieldnames, rows),
        row_count=len(rows),
        nulls=f"{nulls} empty value cells" if nulls else "none",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def _is_value_number_row(row: dict[str, str], value_names: tuple[str, ...]) -> bool:
    values = [row[name] for name in value_names if row[name]]
    if not values:
        return False
    return all(value.isdigit() and len(value) <= 2 for value in values)


def _page_head(words: list[PdfWord], count: int = 20) -> str:
    return " ".join(
        word.text
        for word in sorted(words, key=lambda item: (item.top, item.x0))[:count]
    )


def _collapsed(text: str) -> str:
    return text.replace(" ", "").replace("–", "-").replace("—", "-")


def _open_pdf(payload: bytes, label: str) -> tuple[Any, ParsedTable | None]:
    if not payload.startswith(b"%PDF"):
        return None, _empty_fail(f"{label} is not a pdf")
    try:
        return pdfplumber.open(BytesIO(payload)), None
    except PdfminerException:
        return None, _empty_fail(f"{label}: not a readable pdf")


def _group_metrics(assigned: AssignedTable) -> tuple[tuple[float, ...], float]:
    centers = group_centers(assigned.centers, GROUP_WIDTH)
    if len(centers) < 2:
        return centers, assigned.half
    gaps = [centers[i + 1] - centers[i] for i in range(len(centers) - 1)]
    return centers, min(gaps) / 2


def _assign_grouped(
    words: list[PdfWord],
    make_columns: Callable[[int], tuple[PdfColumn, ...]],
    group_counts: tuple[int, ...],
) -> AssignedTable | None:
    for count in group_counts:
        assigned = assign_pdf_table(words, make_columns(count), y_tol=PDF_Y_TOL)
        if assigned is None:
            continue
        return assigned
    return None


def _fertility_pdf_columns(n_groups: int) -> tuple[PdfColumn, ...]:
    return tuple(
        PdfColumn(f"g{group}_{resid}", resid)
        for group in range(n_groups)
        for resid in ("Total", "Rural", "Urban")
    )


def _table8_pdf_columns(n_groups: int) -> tuple[PdfColumn, ...]:
    return tuple(
        PdfColumn(f"g{group}_{sex}", token)
        for group in range(n_groups)
        for sex, token in (
            ("persons", "Persons"),
            ("male", "Male"),
            ("females", "Females"),
        )
    )


def _label_bucket(
    buckets: list[list[PdfWord]], assigned: AssignedTable, label: str
) -> list[PdfWord] | None:
    hits = [
        bucket
        for bucket in buckets[assigned.header_index + 1 :]
        if " ".join(
            word.text
            for word in bucket
            if word.xmid < assigned.left or not is_amount_token(word.text)
        )
        == label
    ]
    if len(hits) != 1:
        return None
    return hits[0]


def _stitch_and_lines(rows: list[dict[str, str]]) -> list[dict[str, str]] | None:
    out: list[dict[str, str]] = []
    for row in rows:
        label = row["line_label"]
        values = [row[name] for name in BULLETIN_VALUE_COLUMNS if row[name]]
        if label.startswith("and ") and not values:
            if not out:
                return None
            previous = dict(out[-1])
            previous["line_label"] = f"{previous['line_label']} {label}"
            out[-1] = previous
            continue
        out.append(row)
    return out


def _unpivot_table3(
    words: list[PdfWord], assigned: AssignedTable
) -> list[dict[str, str]] | None:
    n_groups = len(assigned.centers) // GROUP_WIDTH
    centers, half = _group_metrics(assigned)
    buckets = cluster_rows(words, PDF_Y_TOL)
    current: list[str] | None = None
    out: list[dict[str, str]] = []
    value_names = tuple(
        f"g{group}_{resid}"
        for group in range(n_groups)
        for resid in ("Total", "Rural", "Urban")
    )
    for row in assigned.rows:
        if _is_value_number_row(row, value_names):
            continue
        has_values = any(row[name] for name in value_names)
        label = row["line_label"]
        if not has_values:
            if label.isdigit():
                continue
            bucket = _label_bucket(buckets, assigned, label)
            if bucket is None:
                return None
            names = names_from_groups(bucket, centers, half, left=assigned.left)
            if names is None and n_groups == 1 and label:
                current = [label]
            elif names is None:
                return None
            else:
                current = names
            continue
        if current is None:
            return None
        for group, geography in enumerate(current):
            out.append(
                {
                    "geography": geography,
                    "line_label": label,
                    "total": row[f"g{group}_Total"],
                    "rural": row[f"g{group}_Rural"],
                    "urban": row[f"g{group}_Urban"],
                }
            )
    return out


def _table8_name_words(
    buckets: list[list[PdfWord]], header_index: int
) -> list[PdfWord]:
    collected: list[PdfWord] = []
    for bucket in buckets[:header_index]:
        if any(
            word.text in TABLE8_SKIP_WORDS or "Contd" in word.text for word in bucket
        ):
            continue
        collected.extend(bucket)
    return collected


def _unpivot_table8(
    words: list[PdfWord], assigned: AssignedTable
) -> list[dict[str, str]] | None:
    n_groups = len(assigned.centers) // GROUP_WIDTH
    centers, half = _group_metrics(assigned)
    buckets = cluster_rows(words, PDF_Y_TOL)
    names = names_from_groups(
        _table8_name_words(buckets, assigned.header_index),
        centers,
        half,
        left=assigned.left,
    )
    if names is None:
        return None
    value_names = tuple(
        f"g{group}_{sex}"
        for group in range(n_groups)
        for sex in ("persons", "male", "females")
    )
    out: list[dict[str, str]] = []
    for row in assigned.rows:
        if _is_value_number_row(row, value_names):
            continue
        year = row["line_label"]
        if not year.isdigit() or len(year) != 4:
            continue
        for group, geography in enumerate(names):
            out.append(
                {
                    "geography": geography,
                    "year": year,
                    "persons": row[f"g{group}_persons"],
                    "male": row[f"g{group}_male"],
                    "females": row[f"g{group}_females"],
                }
            )
    return out


def _is_column_number_row(row: tuple[str, ...]) -> bool:
    values = [cell.strip() for cell in row if cell.strip()]
    if not values:
        return False
    return all(value.isdigit() and len(value) <= 2 for value in values)


def _is_census_year(text: str) -> bool:
    return bool(_YEAR.search(text.strip()))


def _join_headers(
    rows: tuple[tuple[str, ...], ...], start: int, depth: int = 3
) -> list[str]:
    end = min(len(rows), start + depth)
    width = max((len(row) for row in rows[start:end]), default=0)
    joined = [""] * width
    for row in rows[start:end]:
        padded = row + ("",) * (width - len(row))
        for index, cell in enumerate(padded[:width]):
            token = cell.strip()
            if not token:
                continue
            existing = joined[index].split()
            if token in existing:
                continue
            joined[index] = (
                f"{joined[index]} {token}".strip() if joined[index] else token
            )
    return joined


def _unique_col(headers: list[str], predicate: Callable[[str], bool]) -> int | None:
    hits = [index for index, header in enumerate(headers) if predicate(header)]
    if len(hits) != 1:
        return None
    return hits[0]


def _a02_columns(headers: list[str]) -> dict[str, int] | None:
    mapping = {
        "state_code": _unique_col(
            headers,
            lambda header: (
                header == "State Code"
                or ("State" in header and "Code" in header and "District" not in header)
            ),
        ),
        "district_code": _unique_col(
            headers, lambda header: "District" in header and "Code" in header
        ),
        "name": _unique_col(
            headers,
            lambda header: "Union Territory" in header or "India/State" in header,
        ),
        "census_year": _unique_col(
            headers, lambda header: "Census" in header and "Year" in header
        ),
        "persons": _unique_col(headers, lambda header: header == "Persons"),
        "variation_absolute": _unique_col(headers, lambda header: "Absolute" in header),
        "variation_percentage": _unique_col(
            headers,
            lambda header: header == "Percentage" or header.endswith(" Percentage"),
        ),
        "males": _unique_col(headers, lambda header: header == "Males"),
        "females": _unique_col(headers, lambda header: header == "Females"),
    }
    if any(index is None for index in mapping.values()):
        return None
    return {name: index for name, index in mapping.items() if index is not None}


def _header_row(rows: tuple[tuple[str, ...], ...]) -> int | None:
    hits = [
        index
        for index, row in enumerate(rows)
        if "Persons" in row and "Males" in row and "Females" in row
    ]
    if len(hits) != 1:
        return None
    return hits[0]


def _pick_a02_sheet(sheets: tuple[XlsSheet, ...]) -> XlsSheet | None:
    named = [sheet for sheet in sheets if sheet.name == A02_SHEET]
    if len(named) == 1:
        return named[0]
    matches = [sheet for sheet in sheets if _header_row(sheet.rows) is not None]
    if len(matches) != 1:
        return None
    return matches[0]


def _get(row: tuple[str, ...], index: int) -> str:
    if index >= len(row):
        return ""
    return row[index]


def parse_census_2011_pca_sd(payload: bytes) -> ParsedTable:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            return _empty_fail("pca workbook has no Sheet1")
        sheet = workbook["Sheet1"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return _empty_fail("pca Sheet1 is empty")
        names = [str(cell) if cell is not None else "" for cell in header]
        index = {name: i for i, name in enumerate(names)}
        missing = [name for name in PCA_COLUMNS if name not in index]
        if missing:
            return _empty_fail("pca Sheet1 missing columns: " + ", ".join(missing))
        out_rows: list[dict[str, str]] = []
        nulls = 0
        parked_district = 0
        for raw in rows:
            level = raw[index["Level"]]
            if level == "DISTRICT":
                parked_district += 1
                continue
            if level not in PCA_LEVELS:
                continue
            record = {
                name: "" if raw[index[name]] is None else str(raw[index[name]]).strip()
                for name in PCA_COLUMNS
            }
            nulls += sum(1 for value in record.values() if value == "")
            out_rows.append(record)
    finally:
        workbook.close()
    if not out_rows:
        return _empty_fail("pca Sheet1 has no India/STATE rows")
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=PCA_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(out_rows)
    flags = f"parked_district_rows={parked_district}; Level in India/STATE only"
    return ParsedTable(
        csv_text=buffer.getvalue(),
        row_count=len(out_rows),
        nulls=f"{nulls} empty cells" if nulls else "none",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def parse_census_2011_a02_xls(payload: bytes) -> ParsedTable:
    if not is_xls_ole(payload):
        return _empty_fail("a02 artifact is not an ole xls")
    sheets = read_xls_sheets(payload)
    if sheets is None:
        return _empty_fail(
            "ole xls is not a readable workbook; not substituting the companion PDF"
        )
    sheet = _pick_a02_sheet(sheets)
    if sheet is None:
        return _empty_fail("a02 workbook has no unique Table A-2 header")
    header_index = _header_row(sheet.rows)
    if header_index is None:
        return _empty_fail("a02 sheet has no Persons/Males/Females header")
    columns = _a02_columns(_join_headers(sheet.rows, header_index))
    if columns is None:
        return _empty_fail(
            "a02 header is missing a unique State/District/Year/Persons column"
        )
    parked_district = 0
    park_block = False
    out_rows: list[dict[str, str]] = []
    for raw in sheet.rows[header_index + 1 :]:
        if _is_column_number_row(raw):
            continue
        year = _get(raw, columns["census_year"]).strip()
        if not _is_census_year(year):
            continue
        district = _cell(_get(raw, columns["district_code"]))
        state_code = _get(raw, columns["state_code"]).strip()
        name = _get(raw, columns["name"]).strip()
        if state_code or name or district:
            park_block = district not in {"", "000"}
        if park_block:
            parked_district += 1
            continue
        record = {
            "state_code": state_code,
            "district_code": district,
            "name": name,
            "census_year": year,
            "persons": _cell(_get(raw, columns["persons"])),
            "variation_absolute": _cell(_get(raw, columns["variation_absolute"])),
            "variation_percentage": _cell(_get(raw, columns["variation_percentage"])),
            "males": _cell(_get(raw, columns["males"])),
            "females": _cell(_get(raw, columns["females"])),
        }
        out_rows.append(record)
    if not out_rows:
        return _empty_fail("a02 sheet has no India/State/UT census-year rows")
    if not any(row["state_code"] == "00" or row["name"] == "INDIA" for row in out_rows):
        return _empty_fail("a02 sheet missing INDIA / state_code 00 row")
    nulls = sum(
        1 for row in out_rows for column in A02_VALUE_COLUMNS if row[column] == ""
    )
    flags = (
        f"xlrd cells; sheet={sheet.name}; district_code=000 only; "
        f"parked_district_rows={parked_district}"
    )
    return ParsedTable(
        csv_text=_to_csv(A02_COLUMNS, out_rows),
        row_count=len(out_rows),
        nulls=f"{nulls} empty value cells",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def parse_srs_bulletin_2024(payload: bytes) -> ParsedTable:
    document, failed = _open_pdf(payload, "srs bulletin")
    if failed is not None:
        return failed
    assert document is not None
    try:
        parts: list[AssignedTable] = []
        for page in document.pages:
            words = words_before(page_words(page), "Note")
            if "Table 1:" not in _page_head(words):
                continue
            assigned = assign_pdf_table(words, BULLETIN_PDF_COLUMNS, y_tol=PDF_Y_TOL)
            if assigned is None:
                continue
            parts.append(assigned)
    finally:
        document.close()
    if not parts:
        return _empty_fail("srs bulletin Table 1 header not found")
    ambiguous = sum(part.ambiguous for part in parts)
    if ambiguous:
        return _empty_fail(
            f"srs bulletin Table 1: {ambiguous} amount tokens did not map uniquely; "
            "not guessing"
        )
    rows: list[dict[str, str]] = []
    for part in parts:
        rows.extend(
            row
            for row in part.rows
            if not _is_value_number_row(row, BULLETIN_VALUE_COLUMNS)
        )
    stitched = _stitch_and_lines(rows)
    if stitched is None:
        return _empty_fail("srs bulletin Table 1 wrapped geography line did not stitch")
    ordered = [
        {name: row.get(name, "") for name in BULLETIN_COLUMNS} for row in stitched
    ]
    if not any(row["line_label"] == "India" for row in ordered):
        return _empty_fail("srs bulletin Table 1 missing India row")
    if not any(
        "Dadra & Nagar Haveli and Daman & Diu" in row["line_label"] for row in ordered
    ):
        return _empty_fail(
            "srs bulletin Table 1 missing merged DNH and Daman & Diu row"
        )
    assigned_n = sum(part.assigned for part in parts)
    return _ok_rows(
        BULLETIN_COLUMNS,
        ordered,
        BULLETIN_VALUE_COLUMNS,
        f"pdfplumber words; assigned={assigned_n}; ambiguous=0; y_tol={PDF_Y_TOL}",
    )


def parse_srs_statistical_report_2024(payload: bytes) -> ParsedTable:
    document, failed = _open_pdf(payload, "srs statistical report")
    if failed is not None:
        return failed
    assert document is not None
    try:
        found = False
        assigned_n = 0
        rows: list[dict[str, str]] = []
        for page in document.pages:
            words = page_words(page)
            head = _page_head(words, 12)
            if found and head.startswith("Table 4"):
                break
            if not (head.startswith("Table 3") and "Fertility" in head):
                continue
            found = True
            assigned = _assign_grouped(words, _fertility_pdf_columns, (2, 1))
            if assigned is None:
                return _empty_fail(
                    "srs statistical report Table 3: Total/Rural/Urban header not found"
                )
            if assigned.ambiguous:
                return _empty_fail(
                    f"srs statistical report Table 3: {assigned.ambiguous} amount "
                    "tokens did not map uniquely; not guessing"
                )
            part = _unpivot_table3(words, assigned)
            if part is None:
                return _empty_fail(
                    "srs statistical report Table 3 geography names did not map "
                    "uniquely to column groups; not guessing"
                )
            assigned_n += assigned.assigned
            rows.extend(part)
    finally:
        document.close()
    if not rows:
        return _empty_fail(
            "srs statistical report Detailed Table 3 Fertility Indicators not found"
        )
    if not any(row["geography"] == "India" for row in rows):
        return _empty_fail("srs statistical report Table 3 missing India row")
    if not any(row["line_label"] == "Total Fertility Rate" for row in rows):
        return _empty_fail("srs statistical report Table 3 missing TFR row")
    return _ok_rows(
        TABLE3_COLUMNS,
        rows,
        TABLE3_VALUE_COLUMNS,
        f"pdfplumber words; assigned={assigned_n}; ambiguous=0; y_tol={PDF_Y_TOL}",
    )


def parse_ncp_projections_table8(payload: bytes) -> ParsedTable:
    if not payload.startswith(b"%PDF"):
        return _empty_fail("projection report is not a pdf")
    image_count = payload.count(b"/Image")
    document, failed = _open_pdf(payload, "projection report")
    if failed is not None:
        if image_count > 0:
            return _empty_fail(
                f"scanned pages (Image objects={image_count}); Table 8 cells are not "
                "machine-readable; do not OCR a blended stock"
            )
        return failed
    assert document is not None
    try:
        found = False
        assigned_n = 0
        rows: list[dict[str, str]] = []
        for page in document.pages:
            words = page_words(page)
            collapsed = _collapsed(_page_head(words, 16))
            if found and (
                "TABLE-9" in collapsed
                or "TABLE-10" in collapsed
                or "TABLE-11" in collapsed
            ):
                break
            if "TABLE-8" not in collapsed:
                continue
            found = True
            clipped = words_through(words, "2036")
            assigned = _assign_grouped(clipped, _table8_pdf_columns, (3, 2))
            if assigned is None:
                return _empty_fail(
                    "projection Table 8: Persons/Male/Females header not found"
                )
            if assigned.ambiguous:
                return _empty_fail(
                    f"projection Table 8: {assigned.ambiguous} amount tokens did not "
                    "map uniquely; not guessing"
                )
            part = _unpivot_table8(clipped, assigned)
            if part is None:
                return _empty_fail(
                    "projection Table 8 geography names did not map uniquely to "
                    "column groups; not guessing"
                )
            assigned_n += assigned.assigned
            rows.extend(part)
    finally:
        document.close()
    if not rows:
        if image_count > 0:
            return _empty_fail(
                f"scanned pages (Image objects={image_count}); Table 8 cells are not "
                "machine-readable; do not OCR a blended stock"
            )
        return _empty_fail(
            "Table 8 cells are not machine-readable without guessing; stop"
        )
    if not any(row["geography"] == "INDIA" for row in rows):
        return _empty_fail("projection Table 8 missing INDIA rows")
    years = {row["year"] for row in rows}
    if "2011" not in years or "2036" not in years:
        return _empty_fail("projection Table 8 missing 2011-2036 year rows")
    return _ok_rows(
        TABLE8_COLUMNS,
        rows,
        TABLE8_VALUE_COLUMNS,
        f"pdfplumber words; assigned={assigned_n}; ambiguous=0; y_tol={PDF_Y_TOL}",
    )
