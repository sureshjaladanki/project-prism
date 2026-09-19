"""Tidy C3 producer tables. Stop when the artifact cannot be read as cells."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO

import pdfplumber  # type: ignore[import-untyped]
from openpyxl import load_workbook
from pdfplumber.utils.exceptions import (
    PdfminerException,  # type: ignore[import-untyped]
)

from prism.ingest.parse import ParsedTable
from prism.ingest.parse_c3_layout import (
    AssignedTable,
    PdfColumn,
    PdfWord,
    assign_pdf_table,
    expand_html_table,
    parse_html_tables,
)
from prism.schema import YesNo

PARSER_NAME = "union-c3"
PARSER_VERSION = "1.2.0"
PARSER = f"{PARSER_NAME}-{PARSER_VERSION}"

RECEIPT_SHEET = "ReceiptReport123"
STAT1_SHEET = "Statement1"
DEFICIT_SHEET = "Deficit Statistics"

RECEIPT_COLUMNS = (
    "line_label",
    "major_head",
    "actuals_2024_2025",
    "budget_2025_2026",
    "revised_2025_2026",
    "budget_2026_2027",
)
STAT1_COLUMNS = (
    "line_label",
    "actuals_2024_2025_revenue",
    "actuals_2024_2025_capital",
    "actuals_2024_2025_total",
    "budget_2025_2026_revenue",
    "budget_2025_2026_capital",
    "budget_2025_2026_total",
    "revised_2025_2026_revenue",
    "revised_2025_2026_capital",
    "revised_2025_2026_total",
    "budget_2026_2027_revenue",
    "budget_2026_2027_capital",
    "budget_2026_2027_total",
)
DEFICIT_COLUMNS = (
    "line_label",
    "actuals_2024_2025",
    "budget_2025_2026",
    "revised_2025_2026",
    "budget_2026_2027",
)
ANNEX1_COLUMNS = (
    "line_label",
    "actual_2017_18",
    "actual_2018_19",
    "actual_2019_20",
    "actual_2020_21",
    "actual_2021_22",
    "actual_2022_23",
    "actual_2023_24",
    "actual_2024_25",
    "re_2025_26",
    "be_2026_27",
)
ANNEX1_PDF_COLUMNS = (
    PdfColumn("actual_2017_18", "2017-18"),
    PdfColumn("actual_2018_19", "2018-19"),
    PdfColumn("actual_2019_20", "2019-20"),
    PdfColumn("actual_2020_21", "2020-21"),
    PdfColumn("actual_2021_22", "2021-22"),
    PdfColumn("actual_2022_23", "2022-23"),
    PdfColumn("actual_2023_24", "2023-24"),
    PdfColumn("actual_2024_25", "2024-25"),
    PdfColumn("re_2025_26", "2025-26"),
    PdfColumn("be_2026_27", "2026-27"),
)
LIABILITIES_COLUMNS = (
    "line_label",
    "end_1950_51",
    "end_2021_2022",
    "end_2022_2023",
    "end_2023_24",
    "end_2024_25",
    "revised_2025_26",
    "budget_2026_27",
)
LIABILITIES_PDF_COLUMNS = (
    PdfColumn("end_1950_51", "1950-51"),
    PdfColumn("end_2021_2022", "2021-2022"),
    PdfColumn("end_2022_2023", "2022-2023"),
    PdfColumn("end_2023_24", "2023-24"),
    PdfColumn("end_2024_25", "2024-25"),
    PdfColumn("revised_2025_26", "2025-26"),
    PdfColumn("budget_2026_27", "2026-27"),
)
AFS_COLUMNS = (
    "line_label",
    "actuals_2024_2025",
    "budget_2025_2026",
    "revised_2025_2026",
    "budget_2026_2027",
)
AFS_PDF_COLUMNS = (
    PdfColumn("actuals_2024_2025", "2024-2025"),
    PdfColumn("budget_2025_2026", "2025-2026"),
    PdfColumn("revised_2025_2026", "2025-2026"),
    PdfColumn("budget_2026_2027", "2026-2027"),
)
FA_COLUMNS = (
    "line_label",
    "actuals_2024_2025",
    "actuals_2023_2024",
)
FA_PDF_COLUMNS = (
    PdfColumn("actuals_2024_2025", "2024-2025"),
    PdfColumn("actuals_2023_2024", "2023-2024"),
)
CGA_MONTHLY_COLUMNS = (
    "sl_no",
    "line_label",
    "note",
    "be_2026_2027",
    "actuals_upto_july_2026",
    "pct_current",
    "pct_coppy",
)
_WITHHELD = frozenset({"-", "–", "—", "..", "...", "…", ".", "n.a.", "na", "n/a"})


def _empty_fail(flags: str) -> ParsedTable:
    return ParsedTable(
        csv_text="",
        row_count=0,
        nulls="not parsed",
        flags=flags,
        lineage_ok=YesNo.no,
    )


def _norm(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _cell(value: object) -> str:
    text = _norm(value)
    if text.lower() in _WITHHELD:
        return ""
    return text


def _split_value_indexes(split: list[str], starts: list[int]) -> list[int] | None:
    """Revenue / Capital / Total after each year, skipping a blank before Total."""

    bounds = [*starts, len(split)]
    indexes: list[int] = []
    wanted = ("Revenue", "Capital", "Total")
    for i, start in enumerate(starts):
        found: list[int] = []
        want = 0
        for index in range(start, bounds[i + 1]):
            if want < 3 and split[index] == wanted[want]:
                found.append(index)
                want += 1
        if len(found) != 3:
            return None
        indexes.extend(found)
    return indexes


def _status_matches(cell: str, needle: str) -> bool:
    return needle in cell


def _to_csv(fieldnames: tuple[str, ...], rows: list[dict[str, str]]) -> str:
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def _ok(
    fieldnames: tuple[str, ...],
    rows: list[dict[str, str]],
    value_columns: tuple[str, ...],
    flags: str,
) -> ParsedTable:
    nulls = sum(1 for row in rows for column in value_columns if row[column] == "")
    return ParsedTable(
        csv_text=_to_csv(fieldnames, rows),
        row_count=len(rows),
        nulls=f"{nulls} empty value cells",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def parse_frbm_statements_pdf(payload: bytes) -> ParsedTable:
    if not payload.startswith(b"%PDF"):
        return _empty_fail("budget artifact is not a pdf")
    return _empty_fail(
        "frbm1.pdf is statutory FRBM prose, not a reconstructable receipts/"
        "expenditure/deficit/debt grid; do not ingest page 9 Economic "
        "Performance at a Glance as this card"
    )


def parse_cga_monthly_html(payload: bytes) -> ParsedTable:
    head = payload.lstrip()[:2048].lower()
    if b"<html" not in head and b"<!doctype" not in head:
        return _empty_fail("cga monthly artifact is not html")
    html = payload.decode("cp1252")
    tables = parse_html_tables(html)
    if not tables:
        return _empty_fail("cga monthly html has no table")
    grid = expand_html_table(tables[0])
    if not grid or len(grid[0]) != 7:
        return _empty_fail(
            "cga glance table is not 7 columns after colspan/rowspan expand"
        )
    joined_header = " ".join(cell for row in grid[:2] for cell in row)
    if "Budget Estimates" not in joined_header or "July 2026" not in joined_header:
        return _empty_fail("cga glance header missing Budget Estimates / July 2026")
    out_rows: list[dict[str, str]] = []
    for raw in grid[2:]:
        if len(raw) != 7:
            return _empty_fail("cga glance body row is not 7 cells")
        record = {
            "sl_no": raw[0],
            "line_label": raw[1],
            "note": raw[2],
            "be_2026_2027": raw[3],
            "actuals_upto_july_2026": raw[4],
            "pct_current": raw[5],
            "pct_coppy": raw[6],
        }
        if not any(record.values()):
            continue
        out_rows.append(record)
    labels = {row["line_label"] for row in out_rows}
    if "Revenue Receipts" not in labels:
        return _empty_fail("cga glance missing Revenue Receipts row")
    return _ok(
        CGA_MONTHLY_COLUMNS,
        out_rows,
        CGA_MONTHLY_COLUMNS[3:],
        "html table 0; colspan/rowspan expanded; 7 columns",
    )


def parse_cga_finance_accounts_pdf(payload: bytes) -> ParsedTable:
    return _parse_pdf_word_table(
        payload,
        FA_PDF_COLUMNS,
        FA_COLUMNS,
        "finance accounts Statement No. 1",
        "Central Goods and Services Tax",
    )


def _page_words(page: object) -> list[PdfWord]:
    extract_words = page.extract_words  # type: ignore[attr-defined]
    words: list[PdfWord] = []
    for raw in extract_words() or []:
        words.append(
            PdfWord(
                text=str(raw["text"]),
                x0=float(raw["x0"]),
                x1=float(raw["x1"]),
                top=float(raw["top"]),
            )
        )
    return words


def _is_column_number_row(row: dict[str, str], value_names: tuple[str, ...]) -> bool:
    values = [row[name] for name in value_names if row[name]]
    if not values:
        return False
    return all(value.isdigit() and len(value) <= 2 for value in values)


def _parse_pdf_word_table(
    payload: bytes,
    columns: tuple[PdfColumn, ...],
    fieldnames: tuple[str, ...],
    label: str,
    required_label: str,
) -> ParsedTable:
    if not payload.startswith(b"%PDF"):
        return _empty_fail(f"{label} artifact is not a pdf")
    try:
        document = pdfplumber.open(BytesIO(payload))
    except PdfminerException:
        return _empty_fail(f"{label}: not a readable pdf")
    try:
        parts: list[AssignedTable] = []
        for page in document.pages:
            assigned = assign_pdf_table(_page_words(page), columns)
            if assigned is None:
                continue
            parts.append(assigned)
    finally:
        document.close()
    if not parts:
        return _empty_fail(f"{label}: year header not found")
    ambiguous = sum(part.ambiguous for part in parts)
    assigned_n = sum(part.assigned for part in parts)
    rows: list[dict[str, str]] = []
    for part in parts:
        rows.extend(part.rows)
    if ambiguous:
        return _empty_fail(
            f"{label}: {ambiguous} amount tokens did not map uniquely "
            "to a year column; not guessing"
        )
    if not rows:
        return _empty_fail(f"{label}: no amount rows")
    ordered = [
        {name: row.get(name, "") for name in fieldnames}
        for row in rows
        if not _is_column_number_row(row, fieldnames[1:])
    ]
    if not any(row.get(name) for row in ordered for name in fieldnames[1:]):
        return _empty_fail(f"{label}: no amount rows")
    if not any(required_label in row["line_label"] for row in ordered):
        return _empty_fail(f"{label}: missing {required_label!r} row")
    return _ok(
        fieldnames,
        ordered,
        fieldnames[1:],
        f"pdfplumber words; assigned={assigned_n}; ambiguous=0",
    )


def parse_annex1_pdf(payload: bytes) -> ParsedTable:
    return _parse_pdf_word_table(
        payload,
        ANNEX1_PDF_COLUMNS,
        ANNEX1_COLUMNS,
        "annex-1 trends in receipts",
        "REVENUE RECEIPTS",
    )


def parse_liabilities_pdf(payload: bytes) -> ParsedTable:
    return _parse_pdf_word_table(
        payload,
        LIABILITIES_PDF_COLUMNS,
        LIABILITIES_COLUMNS,
        "statement of liabilities",
        "Public Debt",
    )


def parse_afs_pdf(payload: bytes) -> ParsedTable:
    return _parse_pdf_word_table(
        payload,
        AFS_PDF_COLUMNS,
        AFS_COLUMNS,
        "annual financial statement",
        "Corporation Tax",
    )


def parse_receipt_xlsx(payload: bytes) -> ParsedTable:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if RECEIPT_SHEET not in workbook.sheetnames:
            return _empty_fail(f"receipt workbook has no {RECEIPT_SHEET}")
        rows = list(workbook[RECEIPT_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    header_index = None
    major_at = None
    year_at: dict[str, int] = {}
    for i, raw in enumerate(rows):
        cells = [_norm(cell) for cell in raw]
        if "Actuals 2024-2025" not in cells or "Budget 2026-2027" not in cells:
            continue
        if "Budget 2025-2026" not in cells or "Revised 2025-2026" not in cells:
            return _empty_fail("receipt year header missing BE/RE 2025-2026")
        if "Major Head" not in cells:
            return _empty_fail("receipt year header missing Major Head")
        header_index = i
        major_at = cells.index("Major Head")
        year_at = {
            "actuals_2024_2025": cells.index("Actuals 2024-2025"),
            "budget_2025_2026": cells.index("Budget 2025-2026"),
            "revised_2025_2026": cells.index("Revised 2025-2026"),
            "budget_2026_2027": cells.index("Budget 2026-2027"),
        }
        break
    if header_index is None or major_at is None:
        return _empty_fail("receipt sheet has no Actuals 2024-2025 year header")
    out_rows: list[dict[str, str]] = []
    for raw in rows[header_index + 1 :]:
        cells = list(raw)
        label_parts = [
            _cell(cells[i]) for i in range(min(major_at, len(cells))) if _cell(cells[i])
        ]
        record = {
            "line_label": " ".join(label_parts),
            "major_head": _cell(cells[major_at]) if major_at < len(cells) else "",
            "actuals_2024_2025": _cell(cells[year_at["actuals_2024_2025"]])
            if year_at["actuals_2024_2025"] < len(cells)
            else "",
            "budget_2025_2026": _cell(cells[year_at["budget_2025_2026"]])
            if year_at["budget_2025_2026"] < len(cells)
            else "",
            "revised_2025_2026": _cell(cells[year_at["revised_2025_2026"]])
            if year_at["revised_2025_2026"] < len(cells)
            else "",
            "budget_2026_2027": _cell(cells[year_at["budget_2026_2027"]])
            if year_at["budget_2026_2027"] < len(cells)
            else "",
        }
        if not any(record.values()):
            continue
        out_rows.append(record)
    if not out_rows:
        return _empty_fail("receipt sheet has no data rows")
    return _ok(
        RECEIPT_COLUMNS,
        out_rows,
        RECEIPT_COLUMNS[2:],
        f"sheet={RECEIPT_SHEET}",
    )


def parse_expenditure_stat1_xlsx(payload: bytes) -> ParsedTable:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if STAT1_SHEET not in workbook.sheetnames:
            return _empty_fail(f"expenditure workbook has no {STAT1_SHEET}")
        rows = list(workbook[STAT1_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    year_row_index = None
    starts: list[int] = []
    for i, raw in enumerate(rows):
        cells = [_norm(cell) for cell in raw]
        if "Actuals 2024-2025" not in cells:
            continue
        if "Budget Estimates 2026-2027" not in cells:
            return _empty_fail(
                "statement 1 year header missing Budget Estimates 2026-2027"
            )
        needed = (
            "Actuals 2024-2025",
            "Budget Estimates 2025-2026",
            "Revised Estimates 2025-2026",
            "Budget Estimates 2026-2027",
        )
        missing = [name for name in needed if name not in cells]
        if missing:
            return _empty_fail("statement 1 year header missing: " + ", ".join(missing))
        year_row_index = i
        starts = [cells.index(name) for name in needed]
        break
    if year_row_index is None or year_row_index + 1 >= len(rows):
        return _empty_fail("statement 1 has no Actuals 2024-2025 year header")
    split = [_norm(cell) for cell in rows[year_row_index + 1]]
    value_indexes = _split_value_indexes(split, starts)
    if value_indexes is None:
        return _empty_fail(
            "statement 1 split header is not Revenue/Capital/Total at each year"
        )
    out_rows: list[dict[str, str]] = []
    label_end = min(starts)
    for raw in rows[year_row_index + 2 :]:
        cells = list(raw)
        label = " ".join(
            _cell(cells[i])
            for i in range(min(label_end, len(cells)))
            if _cell(cells[i])
        )
        values: list[str] = []
        for index in value_indexes:
            values.append(_cell(cells[index]) if index < len(cells) else "")
        record = {"line_label": label}
        for name, value in zip(STAT1_COLUMNS[1:], values, strict=True):
            record[name] = value
        if not any(record.values()):
            continue
        out_rows.append(record)
    if not out_rows:
        return _empty_fail("statement 1 has no data rows")
    return _ok(STAT1_COLUMNS, out_rows, STAT1_COLUMNS[1:], f"sheet={STAT1_SHEET}")


def parse_deficit_statistics_xlsx(payload: bytes) -> ParsedTable:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if DEFICIT_SHEET not in workbook.sheetnames:
            return _empty_fail(f"BAG workbook has no {DEFICIT_SHEET!r} sheet")
        rows = list(workbook[DEFICIT_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    year_row_index = None
    year_at: list[int] = []
    for i, raw in enumerate(rows):
        cells = [_norm(cell) for cell in raw]
        hits = [
            index
            for index, cell in enumerate(cells)
            if cell in {"2024-2025", "2025-2026", "2026-2027"}
        ]
        if cells.count("2024-2025") != 1 or cells.count("2026-2027") != 1:
            continue
        if cells.count("2025-2026") != 2:
            continue
        year_row_index = i
        year_at = hits
        break
    if year_row_index is None or len(year_at) != 4 or year_row_index + 1 >= len(rows):
        return _empty_fail("deficit sheet has no Actuals/BE/RE/BE year header")
    status = [_norm(cell) for cell in rows[year_row_index + 1]]
    expected_status = (
        "Actuals",
        "Budget Estimates",
        "Revised Estimates",
        "Budget Estimates",
    )
    got_status = tuple(
        status[index] if index < len(status) else "" for index in year_at
    )
    if not all(
        _status_matches(cell, needle)
        for cell, needle in zip(got_status, expected_status, strict=True)
    ):
        return _empty_fail("deficit status row is not Actuals / BE / RE / BE")
    label_end = min(year_at)
    out_rows: list[dict[str, str]] = []
    for raw in rows[year_row_index + 2 :]:
        cells = list(raw)
        label = " ".join(
            _cell(cells[i])
            for i in range(min(label_end, len(cells)))
            if _cell(cells[i])
        )
        record = {
            "line_label": label,
            "actuals_2024_2025": _cell(cells[year_at[0]])
            if year_at[0] < len(cells)
            else "",
            "budget_2025_2026": _cell(cells[year_at[1]])
            if year_at[1] < len(cells)
            else "",
            "revised_2025_2026": _cell(cells[year_at[2]])
            if year_at[2] < len(cells)
            else "",
            "budget_2026_2027": _cell(cells[year_at[3]])
            if year_at[3] < len(cells)
            else "",
        }
        if not any(record.values()):
            continue
        out_rows.append(record)
    if not out_rows:
        return _empty_fail("deficit sheet has no data rows")
    return _ok(
        DEFICIT_COLUMNS,
        out_rows,
        DEFICIT_COLUMNS[1:],
        f"sheet={DEFICIT_SHEET}",
    )
