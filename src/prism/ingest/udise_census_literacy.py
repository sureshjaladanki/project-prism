"""UDISE+ open-services JSON and Census PCA literacy parsers for C6."""

from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from typing import Any

import httpx
from openpyxl import load_workbook  # type: ignore[import-untyped]

from prism.catalog.registry import ParserSpec, RetrieverSpec, register_parser, register_retriever
from prism.ingest.parsed_table import ParsedTable
from prism.ingest.retrieve import RetrievedArtifact, sha256_hex
from prism.schema import YesNo

PARSER_NAME = "dosel-udise-orgi-c6"
PARSER_VERSION = "1.0.0"
PARSER = f"{PARSER_NAME}-{PARSER_VERSION}"

UDISE_YEAR_ID = "12"
UDISE_API_BASE = "https://api.udiseplus.gov.in/open-services/v1.1"
UDISE_BUNDLE_FILENAME = "udise-plus-2025-26-open-services-year12.json"
UDISE_HEADERS = {
    "Content-Type": "application/json",
    "Origin": "https://dashboard.udiseplus.gov.in",
    "Referer": "https://dashboard.udiseplus.gov.in/",
}

SCHOOLS_ENROL_TEACHERS_COLUMNS = (
    "region_code",
    "region_name",
    "schools",
    "enrolments",
    "teachers",
)
FACILITY_VALUE_COLUMNS = (
    "tot_sch_b_toilet",
    "tot_sch_g_toilet",
    "tot_sch_library",
    "tot_sch_electricity",
    "tot_sch_drinkwater",
    "tot_sch_handwash",
    "tot_sch_medical",
    "tot_sch_ramp",
)
FACILITY_COLUMNS = ("region_code", "region_name") + FACILITY_VALUE_COLUMNS
FACILITY_API_KEYS = (
    "totSchBToilet",
    "totSchGToilet",
    "totSchLibrary",
    "totSchElectricity",
    "totSchDrinkwater",
    "totSchHandwash",
    "totSchMedical",
    "totSchRamp",
)

LITERACY_COLUMNS = (
    "State",
    "District",
    "Level",
    "Name",
    "TRU",
    "TOT_P",
    "TOT_M",
    "TOT_F",
    "P_LIT",
    "M_LIT",
    "F_LIT",
    "P_ILL",
    "M_ILL",
    "F_ILL",
)
PCA_LEVELS = frozenset({"India", "STATE"})


def _empty_fail(flags: str) -> ParsedTable:
    return ParsedTable(
        csv_text="",
        row_count=0,
        nulls="not parsed",
        flags=flags,
        lineage_ok=YesNo.no,
    )


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _post_json(
    client: httpx.Client, path: str, body: dict[str, object]
) -> dict[str, Any]:
    url = f"{UDISE_API_BASE}/{path}"
    response = client.post(url, json=body, headers=UDISE_HEADERS)
    if response.status_code != 200:
        raise ValueError(f"http_{response.status_code} for {url}")
    payload = response.json()
    if payload.get("status") is not True:
        raise ValueError(f"status false for {url}: {payload.get('errorDetails')!r}")
    data = payload.get("data")
    if not isinstance(data, list) or not data:
        raise ValueError(f"empty data for {url}")
    return {"url": url, "body": body, "rows": data}


def retrieve_udise_year12_bundle(
    client: httpx.Client, retrieved_at: str, catalog_url: str
) -> RetrievedArtifact:
    """POST yearId 12 summarised-stats (India + State/UT) into one JSON artifact."""

    states_body: dict[str, object] = {
        "yearId": UDISE_YEAR_ID,
        "regionType": 21,
        "regionCode": 99,
        "valueType": 1,
    }
    india_body: dict[str, object] = {
        "yearId": UDISE_YEAR_ID,
        "regionType": 10,
        "regionCode": 99,
        "valueType": 1,
    }
    try:
        schools_states = _post_json(
            client, "schools-summarised-stats/public", states_body
        )
        schools_india = _post_json(
            client, "schools-summarised-stats/public", india_body
        )
        teachers_states = _post_json(
            client, "teachers-summarised-stats/public", states_body
        )
        teachers_india = _post_json(
            client, "teachers-summarised-stats/public", india_body
        )
        students_states = _post_json(
            client, "students-summarised-stats/public", states_body
        )
        students_india = _post_json(
            client, "students-summarised-stats/public", india_body
        )
    except ValueError as exc:
        # Store a failure envelope so lineage can fail without inventing cells.
        content = json.dumps(
            {"yearId": UDISE_YEAR_ID, "error": str(exc)},
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        ).encode("utf-8") + b"\n"
        return RetrievedArtifact(
            url=catalog_url,
            http_status=200,
            content_type="application/json",
            content=content,
            filename=UDISE_BUNDLE_FILENAME,
            checksum=sha256_hex(content),
            retrieved_at=retrieved_at,
        )

    bundle = {
        "yearId": UDISE_YEAR_ID,
        "catalog_url": catalog_url,
        "schools": {
            "india": schools_india["rows"][0],
            "states": schools_states["rows"],
            "india_url": schools_india["url"],
            "states_url": schools_states["url"],
        },
        "teachers": {
            "india": teachers_india["rows"][0],
            "states": teachers_states["rows"],
            "india_url": teachers_india["url"],
            "states_url": teachers_states["url"],
        },
        "students": {
            "india": students_india["rows"][0],
            "states": students_states["rows"],
            "india_url": students_india["url"],
            "states_url": students_states["url"],
        },
    }
    content = json.dumps(bundle, ensure_ascii=False, indent=2, sort_keys=True).encode(
        "utf-8"
    ) + b"\n"
    return RetrievedArtifact(
        url=catalog_url,
        http_status=200,
        content_type="application/json",
        content=content,
        filename=UDISE_BUNDLE_FILENAME,
        checksum=sha256_hex(content),
        retrieved_at=retrieved_at,
    )


def _load_bundle(payload: bytes) -> dict[str, Any] | None:
    try:
        data = json.loads(payload)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return None
    if not isinstance(data, dict):
        return None
    return data


def _region_code(row: dict[str, Any]) -> str:
    raw = row.get("regionCd", row.get("regionCode", ""))
    text = _cell_str(raw)
    if text.isdigit():
        return text.zfill(2) if len(text) <= 2 else text
    return text


def _index_by_code(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = _region_code(row)
        if code == "":
            continue
        out[code] = row
    return out


def parse_udise_schools_enrolment_teachers(payload: bytes) -> ParsedTable:
    bundle = _load_bundle(payload)
    if bundle is None:
        return _empty_fail("udise bundle is not a json object")
    if "error" in bundle:
        return _empty_fail(f"udise retrieve failed: {bundle['error']}")
    if bundle.get("yearId") != UDISE_YEAR_ID:
        return _empty_fail(
            f"udise yearId {bundle.get('yearId')!r} is not {UDISE_YEAR_ID}"
        )
    try:
        schools = bundle["schools"]
        teachers = bundle["teachers"]
        students = bundle["students"]
        school_india = schools["india"]
        teacher_india = teachers["india"]
        student_india = students["india"]
        school_states = list(schools["states"])
        teacher_by_code = _index_by_code(list(teachers["states"]))
        student_by_code = _index_by_code(list(students["states"]))
    except (KeyError, TypeError) as exc:
        return _empty_fail(f"udise bundle missing schools/teachers/students: {exc}")

    if len(school_states) != 36:
        return _empty_fail(
            f"udise schools states expected 36 rows, got {len(school_states)}"
        )

    out_rows: list[dict[str, str]] = []
    nulls = 0
    flags_parts: list[str] = []

    india_code = _region_code(school_india)
    india_name = _cell_str(school_india.get("regionName"))
    if india_code != "100" or india_name == "":
        return _empty_fail(
            f"udise india row ambiguous code/name {india_code!r}/{india_name!r}"
        )
    # Store the opened artefact string (ALL INDIA on summarised-stats).
    teacher_name = _cell_str(teacher_india.get("regionName"))
    student_name = _cell_str(student_india.get("regionName"))
    if _region_code(teacher_india) != india_code or _region_code(student_india) != india_code:
        return _empty_fail("udise india region codes disagree across endpoints")
    if teacher_name != india_name or student_name != india_name:
        flags_parts.append(
            f"india_name_cross_endpoint={india_name}/{teacher_name}/{student_name}"
        )

    india_record = {
        "region_code": india_code,
        "region_name": india_name,
        "schools": _cell_str(school_india.get("totSchools")),
        "enrolments": _cell_str(student_india.get("totStudents")),
        "teachers": _cell_str(teacher_india.get("totTch")),
    }
    nulls += sum(1 for value in india_record.values() if value == "")
    out_rows.append(india_record)

    for school_row in school_states:
        code = _region_code(school_row)
        name = _cell_str(school_row.get("regionName"))
        if code == "" or name == "":
            return _empty_fail("udise state row missing regionCd/regionName")
        teacher_row = teacher_by_code.get(code)
        student_row = student_by_code.get(code)
        if teacher_row is None or student_row is None:
            return _empty_fail(f"udise missing teachers/students for {code} {name}")
        teacher_label = _cell_str(teacher_row.get("regionName"))
        student_label = _cell_str(student_row.get("regionName"))
        if teacher_label != name or student_label != name:
            return _empty_fail(
                f"udise regionName mismatch for {code}: "
                f"schools={name!r} teachers={teacher_label!r} students={student_label!r}"
            )
        record = {
            "region_code": code,
            "region_name": name,
            "schools": _cell_str(school_row.get("totSchools")),
            "enrolments": _cell_str(student_row.get("totStudents")),
            "teachers": _cell_str(teacher_row.get("totTch")),
        }
        nulls += sum(1 for value in record.values() if value == "")
        out_rows.append(record)

    if len(out_rows) != 37:
        return _empty_fail(f"udise enrolment table expected 37 rows, got {len(out_rows)}")

    buffer = StringIO()
    writer = csv.DictWriter(
        buffer, fieldnames=SCHOOLS_ENROL_TEACHERS_COLUMNS, lineterminator="\n"
    )
    writer.writeheader()
    writer.writerows(out_rows)
    flags = "open-services yearId=12; India+36 State/UT; NEP stocks"
    if flags_parts:
        flags = flags + "; " + "; ".join(flags_parts)
    return ParsedTable(
        csv_text=buffer.getvalue(),
        row_count=len(out_rows),
        nulls=f"{nulls} empty cells" if nulls else "none",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def parse_udise_facilities(payload: bytes) -> ParsedTable:
    bundle = _load_bundle(payload)
    if bundle is None:
        return _empty_fail("udise bundle is not a json object")
    if "error" in bundle:
        return _empty_fail(f"udise retrieve failed: {bundle['error']}")
    if bundle.get("yearId") != UDISE_YEAR_ID:
        return _empty_fail(
            f"udise yearId {bundle.get('yearId')!r} is not {UDISE_YEAR_ID}"
        )
    try:
        schools = bundle["schools"]
        school_india = schools["india"]
        school_states = list(schools["states"])
    except (KeyError, TypeError) as exc:
        return _empty_fail(f"udise bundle missing schools: {exc}")

    if len(school_states) != 36:
        return _empty_fail(
            f"udise schools states expected 36 rows, got {len(school_states)}"
        )

    sample = school_states[0]
    missing_keys = [key for key in FACILITY_API_KEYS if key not in sample]
    if missing_keys:
        return _empty_fail(
            "udise facility columns missing on schools-summarised-stats: "
            + ", ".join(missing_keys)
        )

    out_rows: list[dict[str, str]] = []
    nulls = 0

    def _facility_record(row: dict[str, Any]) -> dict[str, str] | None:
        code = _region_code(row)
        name = _cell_str(row.get("regionName"))
        if code == "" or name == "":
            return None
        record = {"region_code": code, "region_name": name}
        for csv_name, api_key in zip(FACILITY_VALUE_COLUMNS, FACILITY_API_KEYS, strict=True):
            record[csv_name] = _cell_str(row.get(api_key))
        return record

    india = _facility_record(school_india)
    if india is None or india["region_code"] != "100":
        return _empty_fail("udise facilities india row ambiguous")
    nulls += sum(1 for value in india.values() if value == "")
    out_rows.append(india)

    for school_row in school_states:
        record = _facility_record(school_row)
        if record is None:
            return _empty_fail("udise facility state row missing regionCd/regionName")
        nulls += sum(1 for value in record.values() if value == "")
        out_rows.append(record)

    if len(out_rows) != 37:
        return _empty_fail(f"udise facilities expected 37 rows, got {len(out_rows)}")

    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=FACILITY_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(out_rows)
    return ParsedTable(
        csv_text=buffer.getvalue(),
        row_count=len(out_rows),
        nulls=f"{nulls} empty cells" if nulls else "none",
        flags=(
            "open-services yearId=12; having-facility counts from "
            "schools-summarised-stats; not functional columns; not a facilities index"
        ),
        lineage_ok=YesNo.yes,
    )


def parse_census_2011_pca_literacy(payload: bytes) -> ParsedTable:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            return _empty_fail("pca literacy workbook has no Sheet1")
        sheet = workbook["Sheet1"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return _empty_fail("pca literacy Sheet1 is empty")
        names = [str(cell) if cell is not None else "" for cell in header]
        index = {name: i for i, name in enumerate(names)}
        missing = [name for name in LITERACY_COLUMNS if name not in index]
        if missing:
            return _empty_fail(
                "pca literacy Sheet1 missing columns: " + ", ".join(missing)
            )
        out_rows: list[dict[str, str]] = []
        nulls = 0
        parked_district = 0
        for raw in rows:
            level = raw[index["Level"]]
            if level == "DISTRICT":
                parked_district += 1
                continue
            if level not in PCA_LEVELS:
                continue
            record = {
                name: "" if raw[index[name]] is None else str(raw[index[name]]).strip()
                for name in LITERACY_COLUMNS
            }
            nulls += sum(1 for value in record.values() if value == "")
            out_rows.append(record)
    finally:
        workbook.close()
    if not out_rows:
        return _empty_fail("pca literacy Sheet1 has no India/STATE rows")
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=LITERACY_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(out_rows)
    flags = (
        f"parked_district_rows={parked_district}; Level in India/STATE only; "
        "literacy columns only (not C2 population measures)"
    )
    return ParsedTable(
        csv_text=buffer.getvalue(),
        row_count=len(out_rows),
        nulls=f"{nulls} empty cells" if nulls else "none",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


register_retriever(
    "udise-open-services-year12",
    RetrieverSpec(
        retrieve=retrieve_udise_year12_bundle,
        filename=UDISE_BUNDLE_FILENAME,
    ),
)
register_parser(
    "udise-open-services-schools-enrol-teachers",
    ParserSpec(
        parse=parse_udise_schools_enrolment_teachers,
        lineage_name=PARSER,
        kind="json",
    ),
)
register_parser(
    "udise-open-services-facilities",
    ParserSpec(parse=parse_udise_facilities, lineage_name=PARSER, kind="json"),
)
register_parser(
    "census-2011-pca-literacy-xlsx",
    ParserSpec(
        parse=parse_census_2011_pca_literacy, lineage_name=PARSER, kind="xlsx"
    ),
)
