"""Tidy DES / DFPD / FCI foodgrain tables. Stop when cells are ambiguous."""

from __future__ import annotations

import csv
import re
from io import BytesIO, StringIO

import pdfplumber
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pdfplumber.utils.exceptions import PdfminerException

from prism.catalog.registry import ParserSpec, register_parser
from prism.ingest.parsed_table import ParsedTable
from prism.schema import YesNo

PARSER_NAME = "des-dfpd-food-c8"
PARSER_VERSION = "1.0.0"
PARSER = f"{PARSER_NAME}-{PARSER_VERSION}"

AE_FE_COLUMNS = (
    "crop",
    "season",
    "year",
    "production_lakh_tonnes",
    "estimate_round",
)
APY_COLUMNS = (
    "crop",
    "state",
    "season",
    "year",
    "area_thousand_ha",
    "production_thousand_tonnes",
    "yield_kg_ha",
)
PROCUREMENT_COLUMNS = (
    "state",
    "marketing_season",
    "commodity",
    "procurement_lmt",
)
ALLOCATION_OFFTAKE_COLUMNS = (
    "state",
    "period",
    "scheme",
    "allocation_lmt",
    "offtake_lmt",
)
COVERAGE_FPS_COLUMNS = (
    "table",
    "state",
    "measure",
    "value",
)
STOCKS_COLUMNS = (
    "region",
    "rice_lmt",
    "wheat_lmt",
    "total_lmt",
)

_WITHHELD = frozenset(
    {"-", "–", "—", "..", "...", "…", ".", "n.a.", "na", "n/a", "@", "$", "#"}
)
_YEAR = re.compile(r"^(20\d{2})-(\d{2})")
_ZONE_LABELS = (
    "East Zonal Total",
    "North-East Zonal Total",
    "North Zonal Total",
    "South Zonal Total",
    "West Zonal Total",
)


def _empty_fail(flags: str) -> ParsedTable:
    return ParsedTable(
        csv_text="",
        row_count=0,
        nulls="not parsed",
        flags=flags,
        lineage_ok=YesNo.no,
    )


def _to_csv(fieldnames: tuple[str, ...], rows: list[dict[str, str]]) -> str:
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def _norm(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _cell(value: object) -> str:
    text = _norm(value)
    if text.lower() in _WITHHELD:
        return ""
    return text


def _ok(
    fieldnames: tuple[str, ...],
    rows: list[dict[str, str]],
    value_columns: tuple[str, ...],
    flags: str,
) -> ParsedTable:
    nulls = sum(1 for row in rows for name in value_columns if row[name] == "")
    return ParsedTable(
        csv_text=_to_csv(fieldnames, rows),
        row_count=len(rows),
        nulls=f"{nulls} empty value cells" if nulls else "none",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def _english_label(cell: object) -> str:
    """English stub from a bilingual cell. Keep producer spelling; drop Devanagari."""

    text = _norm(cell)
    if not text:
        return ""
    if "/" in text:
        text = text.split("/")[-1].strip()
    parts: list[str] = []
    current: list[str] = []
    for char in text:
        if ord(char) < 128:
            current.append(char)
        else:
            if current:
                parts.append("".join(current).strip(" ,;"))
                current = []
    if current:
        parts.append("".join(current).strip(" ,;"))
    candidates = [part for part in parts if any(char.isalpha() for char in part)]
    if not candidates:
        return text
    best = max(candidates, key=lambda part: (sum(ch.isalpha() for ch in part), len(part)))
    return best.strip().rstrip("(").strip()


def _open_pdf(payload: bytes, label: str):
    if not payload.startswith(b"%PDF"):
        return None, _empty_fail(f"{label} is not a pdf")
    try:
        return pdfplumber.open(BytesIO(payload)), None
    except PdfminerException:
        return None, _empty_fail(f"{label}: not a readable pdf")


def _year_token(header: str) -> str | None:
    collapsed = header.replace("\n", " ")
    match = _YEAR.search(collapsed)
    if match is None:
        return None
    return f"{match.group(1)}-{match.group(2)}"


def _parse_ae_fe_pdf(
    payload: bytes, *, estimate_round: str, required_title: str
) -> ParsedTable:
    document, failed = _open_pdf(payload, "des time-series")
    if failed is not None:
        return failed
    assert document is not None
    try:
        out_rows: list[dict[str, str]] = []
        found_title = False
        for page in document.pages:
            text = page.extract_text() or ""
            if required_title in text.replace("\n", " "):
                found_title = True
            if "Oilseeds" in text and "Food Grains" not in text.split("Oilseeds")[0][-80:]:
                # page 3 of AE/FE is oilseeds/commercial — skip for foodgrain cards
                if "Third Advance Estimate of Production of Oilseeds" in text:
                    continue
                if "Final Estimate of Production of Oilseeds" in text:
                    continue
            tables = page.extract_tables() or []
            if not tables:
                continue
            table = max(tables, key=len)
            if len(table) < 3:
                continue
            year_row = [_norm(cell) for cell in table[1]]
            years: list[tuple[int, str]] = []
            for index, header in enumerate(year_row):
                token = _year_token(header)
                if token is not None:
                    years.append((index, token))
            if len(years) < 5:
                continue
            crop = ""
            for raw in table[2:]:
                cells = [_cell(cell) for cell in raw]
                if not any(cells):
                    continue
                if cells[0] and not cells[0].isdigit() and cells[0] not in {
                    "Crop",
                    "S No.",
                }:
                    # A2 has crop in col 0; A1 has S No then Crop
                    if estimate_round.startswith("3rd"):
                        if len(cells) > 1 and cells[1]:
                            crop = cells[1].replace("\n", " ")
                    else:
                        crop = cells[0].replace("\n", " ")
                elif estimate_round.startswith("3rd") and cells[1]:
                    crop = cells[1].replace("\n", " ")
                elif (not estimate_round.startswith("3rd")) and cells[0]:
                    crop = cells[0].replace("\n", " ")
                season_at = 2 if estimate_round.startswith("3rd") else 1
                season = cells[season_at] if season_at < len(cells) else ""
                if season not in {"Kharif", "Rabi", "Summer", "Total"}:
                    continue
                if not crop:
                    return _empty_fail("des time-series crop label missing before season")
                for index, year in years:
                    value = cells[index] if index < len(cells) else ""
                    out_rows.append(
                        {
                            "crop": crop,
                            "season": season,
                            "year": year,
                            "production_lakh_tonnes": value,
                            "estimate_round": estimate_round,
                        }
                    )
        if not found_title:
            return _empty_fail(f"des time-series missing title {required_title!r}")
        if not out_rows:
            return _empty_fail("des time-series has no foodgrain rows")
        food = [row for row in out_rows if "oil" not in row["crop"].lower()]
        if not any(row["crop"] == "Total Food Grains" for row in food):
            # AE/FE print "Total Food Grains"
            if not any("Total Food" in row["crop"] for row in food):
                return _empty_fail("des time-series missing Total Food Grains row")
        return _ok(
            AE_FE_COLUMNS,
            food,
            ("production_lakh_tonnes",),
            f"pdfplumber extract_tables; estimate_round={estimate_round}",
        )
    finally:
        document.close()


def parse_des_3rd_ae_pdf(payload: bytes) -> ParsedTable:
    return _parse_ae_fe_pdf(
        payload,
        estimate_round="3rd-ae-2025-26",
        required_title="Third Advance Estimate of Production of Food Grains",
    )


def parse_des_fe_pdf(payload: bytes) -> ParsedTable:
    return _parse_ae_fe_pdf(
        payload,
        estimate_round="fe-2024-25",
        required_title="Final Estimate of Production of Food Grains",
    )


def parse_des_apy_xlsx(payload: bytes) -> ParsedTable:
    if not payload.startswith(b"PK\x03\x04"):
        return _empty_fail("des apy artifact is not an xlsx")
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if "Total Foodgrains" not in workbook.sheetnames:
            return _empty_fail("des apy workbook has no Total Foodgrains sheet")
        rows = list(workbook["Total Foodgrains"].iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(rows) < 8:
        return _empty_fail("des apy sheet too short")
    header = [_norm(cell) for cell in rows[6]]
    if header[:3] != ["Crop", "State", "Season"]:
        return _empty_fail("des apy header is not Crop/State/Season")
    years = ("2021-22", "2022-23", "2023-24", "2024-25", "2025-26")
    if header[3:8] != list(years) or header[8:13] != list(years):
        return _empty_fail("des apy year columns are not 2021-22…2025-26 for Area/Production")
    out_rows: list[dict[str, str]] = []
    crop = ""
    state = ""
    for raw in rows[7:]:
        cells = list(raw)
        crop_cell = _cell(cells[0]) if cells else ""
        state_cell = _cell(cells[1]) if len(cells) > 1 else ""
        season = _cell(cells[2]) if len(cells) > 2 else ""
        if crop_cell:
            crop = crop_cell
        if state_cell:
            state = state_cell
        if season not in {"Kharif", "Rabi", "Summer", "Total"}:
            continue
        if not crop or not state:
            return _empty_fail("des apy missing crop/state before season row")
        for offset, year in enumerate(years):
            area = _cell(cells[3 + offset]) if len(cells) > 3 + offset else ""
            production = _cell(cells[8 + offset]) if len(cells) > 8 + offset else ""
            yield_ = _cell(cells[13 + offset]) if len(cells) > 13 + offset else ""
            out_rows.append(
                {
                    "crop": crop,
                    "state": state,
                    "season": season,
                    "year": year,
                    "area_thousand_ha": area,
                    "production_thousand_tonnes": production,
                    "yield_kg_ha": yield_,
                }
            )
    if not out_rows:
        return _empty_fail("des apy sheet has no data rows")
    states = {row["state"] for row in out_rows}
    if "All India" not in states or "J&K" not in states:
        return _empty_fail("des apy missing All India or J&K stub")
    return _ok(
        APY_COLUMNS,
        out_rows,
        ("area_thousand_ha", "production_thousand_tonnes", "yield_kg_ha"),
        "sheet=Total Foodgrains; cell-mapped years 2021-22…2025-26",
    )


def _bulletin_page_by_printed(document: object, printed: str):
    for page in document.pages:  # type: ignore[attr-defined]
        text = page.extract_text() or ""
        for line in text.splitlines()[:4]:
            if line.strip() == printed:
                return page
    return None


def parse_dfpd_bulletin_procurement_pdf(payload: bytes) -> ParsedTable:
    document, failed = _open_pdf(payload, "dfpd bulletin")
    if failed is not None:
        return failed
    assert document is not None
    try:
        page = _bulletin_page_by_printed(document, "9")
        if page is None:
            return _empty_fail("dfpd bulletin missing printed page 9 procurement")
        text = page.extract_text() or ""
        if "Procurement of Rice, Wheat" not in text and "PROCUREMENT" not in text.upper():
            return _empty_fail("dfpd bulletin page 9 is not the procurement table")
        tables = page.extract_tables() or []
        if not tables:
            return _empty_fail("dfpd bulletin page 9 has no table")
        table = max(tables, key=len)
        if len(table) < 3:
            return _empty_fail("dfpd bulletin procurement table too short")
        # Row0: season-year groups; Row1: commodity headers under each year.
        header1 = [_norm(cell) for cell in table[1]]
        # Map column index -> (marketing_season_label, commodity)
        # From probe: cols after state are Rice/Wheat/Coarse for 2022-23…2025-26 then Wheat 2026-27
        seasons = (
            ("2022-23", "Rice", "KMS"),
            ("2022-23", "Wheat", "RMS"),
            ("2022-23", "Coarse grain", "KMS"),
            ("2023-24", "Rice", "KMS"),
            ("2023-24", "Wheat", "RMS"),
            ("2023-24", "Coarse grain", "KMS"),
            ("2024-25", "Rice", "KMS"),
            ("2024-25", "Wheat", "RMS"),
            ("2024-25", "Coarse grain", "KMS"),
            ("2025-26", "Rice", "KMS"),
            ("2025-26", "Wheat", "RMS"),
            ("2025-26", "Coarse grain", "KMS"),
            ("2026-27", "Wheat", "RMS"),
        )
        if len(header1) < 14:
            return _empty_fail("dfpd bulletin procurement header is not 14 columns")
        out_rows: list[dict[str, str]] = []
        for raw in table[2:]:
            cells = [_cell(cell) for cell in raw]
            if not cells or not cells[0]:
                continue
            state = _english_label(raw[0])
            if not state:
                continue
            for offset, (year, commodity, channel) in enumerate(seasons):
                index = offset + 1
                value = cells[index] if index < len(cells) else ""
                out_rows.append(
                    {
                        "state": state,
                        "marketing_season": f"{channel} {year}",
                        "commodity": commodity,
                        "procurement_lmt": value,
                    }
                )
        if not out_rows:
            return _empty_fail("dfpd bulletin procurement has no rows")
        if not any(row["state"] == "Total" for row in out_rows):
            return _empty_fail("dfpd bulletin procurement missing Total row")
        return _ok(
            PROCUREMENT_COLUMNS,
            out_rows,
            ("procurement_lmt",),
            "printed page 9; pdfplumber extract_tables",
        )
    finally:
        document.close()


def parse_dfpd_bulletin_allocation_offtake_pdf(payload: bytes) -> ParsedTable:
    document, failed = _open_pdf(payload, "dfpd bulletin")
    if failed is not None:
        return failed
    assert document is not None
    try:
        # Prefer "upto August" (printed page 16) over single-month (printed page 15).
        page = _bulletin_page_by_printed(document, "16")
        period = "2026-27 upto August 2026"
        if page is None:
            return _empty_fail("dfpd bulletin missing printed page 16 allocation/offtake")
        text = page.extract_text() or ""
        if "Offtake" not in text or "Allocation" not in text:
            return _empty_fail("dfpd bulletin page 16 is not allocation/offtake")
        tables = page.extract_tables() or []
        if not tables:
            return _empty_fail("dfpd bulletin page 16 has no table")
        table = max(tables, key=len)
        # Columns after state: AAY alloc, PHH alloc, Tide alloc, Total alloc,
        # AAY offtake, PHH offtake, Tide offtake, Total offtake, …
        out_rows: list[dict[str, str]] = []
        for raw in table:
            cells = [_cell(cell) for cell in raw]
            if len(cells) < 10:
                continue
            state = _english_label(raw[1] if raw[1] else raw[0])
            if not state or state in {"States/", "States/UTs"}:
                continue
            if not any(ch.isdigit() for ch in "".join(cells[2:10])):
                continue
            schemes = (
                ("AAY", 2, 6),
                ("PHH", 3, 7),
                ("Tide over", 4, 8),
                ("Total", 5, 9),
            )
            for scheme, alloc_at, offtake_at in schemes:
                out_rows.append(
                    {
                        "state": state,
                        "period": period,
                        "scheme": scheme,
                        "allocation_lmt": cells[alloc_at] if alloc_at < len(cells) else "",
                        "offtake_lmt": cells[offtake_at] if offtake_at < len(cells) else "",
                    }
                )
        if not out_rows:
            return _empty_fail("dfpd bulletin allocation/offtake has no rows")
        if not any(row["state"] == "Total" for row in out_rows):
            return _empty_fail("dfpd bulletin allocation/offtake missing Total row")
        return _ok(
            ALLOCATION_OFFTAKE_COLUMNS,
            out_rows,
            ("allocation_lmt", "offtake_lmt"),
            "printed page 16; allocation and offtake kept as separate measures",
        )
    finally:
        document.close()


def parse_dfpd_bulletin_nfsa_coverage_fps_pdf(payload: bytes) -> ParsedTable:
    document, failed = _open_pdf(payload, "dfpd bulletin")
    if failed is not None:
        return failed
    assert document is not None
    try:
        out_rows: list[dict[str, str]] = []
        coverage = _bulletin_page_by_printed(document, "20")
        if coverage is None:
            return _empty_fail("dfpd bulletin missing printed page 20 NFSA coverage")
        tables = coverage.extract_tables() or []
        if not tables:
            return _empty_fail("dfpd bulletin page 20 has no table")
        table = max(tables, key=len)
        for raw in table:
            if not raw or len(raw) < 16:
                continue
            state = _english_label(raw[1] if len(raw) > 1 else "")
            if not state or state.startswith("States") or state in {"UTs"}:
                continue
            serial = _cell(raw[0]).replace(".", "")
            if not serial.isdigit():
                continue
            # Total persons column near end (index 15 in probe)
            total_persons = _cell(raw[15]) if len(raw) > 15 else ""
            aay_hh = _cell(raw[12]) if len(raw) > 12 else ""
            out_rows.append(
                {
                    "table": "nfsa-coverage",
                    "state": state,
                    "measure": "total_persons_lakh",
                    "value": total_persons,
                }
            )
            out_rows.append(
                {
                    "table": "nfsa-coverage",
                    "state": state,
                    "measure": "aay_households_lakh",
                    "value": aay_hh,
                }
            )

        fps = _bulletin_page_by_printed(document, "32")
        if fps is None:
            return _empty_fail("dfpd bulletin missing printed page 32 fair price shops")
        fps_tables = fps.extract_tables() or []
        if not fps_tables:
            return _empty_fail("dfpd bulletin page 32 has no table")
        fps_table = max(fps_tables, key=len)
        for raw in fps_table:
            if not raw or len(raw) < 3:
                continue
            state = _english_label(raw[1])
            if not state or state.startswith("States") or state in {"UTs"}:
                continue
            shops = _cell(raw[2]).replace(",", "")
            out_rows.append(
                {
                    "table": "fair-price-shops",
                    "state": state,
                    "measure": "fair_price_shops",
                    "value": shops,
                }
            )
        if not out_rows:
            return _empty_fail("dfpd bulletin coverage/FPS has no rows")
        return _ok(
            COVERAGE_FPS_COLUMNS,
            out_rows,
            ("value",),
            "printed pages 20 and 32; coverage and FPS as separate measures",
        )
    finally:
        document.close()


def parse_fci_stocks_pdf(payload: bytes) -> ParsedTable:
    document, failed = _open_pdf(payload, "fci stocks")
    if failed is not None:
        return failed
    assert document is not None
    try:
        if not document.pages:
            return _empty_fail("fci stocks pdf has no pages")
        page = document.pages[0]
        text = page.extract_text() or ""
        if "TOTAL STOCKS OF FOODGRAINS IN CENTRAL POOL" not in text:
            return _empty_fail("fci stocks title not found")
        if "01.09.2026" not in text and "01.09.2026" not in text.replace(" ", ""):
            if "AS ON 01.09.2026" not in text:
                return _empty_fail("fci stocks as-on date 01.09.2026 not found")
        tables = page.extract_tables() or []
        if len(tables) < 2:
            return _empty_fail("fci stocks missing main stock table")
        table = tables[1]
        out_rows: list[dict[str, str]] = []
        zone_i = 0
        for raw in table:
            if not raw or not raw[0]:
                continue
            label = _english_label(raw[0])
            if not label or label in {"Region", "RICE", "2"}:
                continue
            if label.startswith("Part") or "Figs" in label:
                continue
            if label == "Zonal Total":
                if zone_i >= len(_ZONE_LABELS):
                    return _empty_fail("fci stocks unexpected extra Zonal Total")
                label = _ZONE_LABELS[zone_i]
                zone_i += 1
            if label.startswith("Wheat lying"):
                label = "Wheat lying in mandies"
            cells = [_cell(cell) for cell in raw]
            # TOTAL CENTRAL POOL STOCK rice/wheat/total = cols 7,8,9 (0-based)
            if len(cells) < 10:
                continue
            if not any(cells[7:10]):
                continue
            out_rows.append(
                {
                    "region": label,
                    "rice_lmt": cells[7],
                    "wheat_lmt": cells[8],
                    "total_lmt": cells[9],
                }
            )
        if zone_i != len(_ZONE_LABELS):
            return _empty_fail(
                f"fci stocks expected {len(_ZONE_LABELS)} zonal totals, got {zone_i}"
            )
        if not any(row["region"] == "Total" for row in out_rows):
            return _empty_fail("fci stocks missing Total row")
        return _ok(
            STOCKS_COLUMNS,
            out_rows,
            ("rice_lmt", "wheat_lmt", "total_lmt"),
            "pdfplumber table 1; zonal totals disambiguated by zone order; "
            "producer=FCI fetch_host=DFPD",
        )
    finally:
        document.close()


def parse_des_asag_pdf(payload: bytes) -> ParsedTable:
    if not payload.startswith(b"%PDF"):
        return _empty_fail("des asag artifact is not a pdf")
    return _empty_fail(
        "ASAG 2024-25 foodgrain tables are not cell-mapped this iteration; "
        "prefer Cards A1–A3 for the live production vintage"
    )


def parse_nfsa_rc_dashboard_html(payload: bytes) -> ParsedTable:
    head = payload.lstrip()[:2048].lower()
    if b"<html" not in head and b"<!doctype" not in head:
        return _empty_fail("nfsa dashboard artifact is not html")
    return _empty_fail(
        "NFSA Public RC Dashboard is live HTML without a verified cell-mapped "
        "export; prefer Cards B1/B2 bulletin tables"
    )


def parse_dfpd_annual_report_pdf(payload: bytes) -> ParsedTable:
    if not payload.startswith(b"%PDF"):
        return _empty_fail("dfpd annual report artifact is not a pdf")
    return _empty_fail(
        "DFPD Annual Report 2025-26 is companion narrative; prefer Cards A5/B1 "
        "for monthly offtake/procurement"
    )


register_parser(
    "des-3rd-ae-pdf",
    ParserSpec(parse=parse_des_3rd_ae_pdf, lineage_name=PARSER, kind="pdf"),
)
register_parser(
    "des-fe-pdf",
    ParserSpec(parse=parse_des_fe_pdf, lineage_name=PARSER, kind="pdf"),
)
register_parser(
    "des-apy-xlsx",
    ParserSpec(parse=parse_des_apy_xlsx, lineage_name=PARSER, kind="xlsx"),
)
register_parser(
    "des-asag-pdf",
    ParserSpec(parse=parse_des_asag_pdf, lineage_name=PARSER, kind="pdf"),
)
register_parser(
    "dfpd-bulletin-procurement-pdf",
    ParserSpec(
        parse=parse_dfpd_bulletin_procurement_pdf, lineage_name=PARSER, kind="pdf"
    ),
)
register_parser(
    "dfpd-bulletin-allocation-offtake-pdf",
    ParserSpec(
        parse=parse_dfpd_bulletin_allocation_offtake_pdf,
        lineage_name=PARSER,
        kind="pdf",
    ),
)
register_parser(
    "dfpd-bulletin-nfsa-coverage-fps-pdf",
    ParserSpec(
        parse=parse_dfpd_bulletin_nfsa_coverage_fps_pdf,
        lineage_name=PARSER,
        kind="pdf",
    ),
)
register_parser(
    "fci-stocks-pdf",
    ParserSpec(parse=parse_fci_stocks_pdf, lineage_name=PARSER, kind="pdf"),
)
register_parser(
    "nfsa-rc-dashboard-html",
    ParserSpec(parse=parse_nfsa_rc_dashboard_html, lineage_name=PARSER, kind="html"),
)
register_parser(
    "dfpd-annual-report-pdf",
    ParserSpec(parse=parse_dfpd_annual_report_pdf, lineage_name=PARSER, kind="pdf"),
)
