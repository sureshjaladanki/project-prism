"""Tidy PLFS Appendix XLSX and Monthly Bulletin Statement tables. No invented cells."""

from __future__ import annotations

import csv
import re
from io import BytesIO, StringIO
from typing import Any

import pdfplumber
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pdfplumber.utils.exceptions import PdfminerException

from prism.catalog.registry import ParserSpec, register_parser
from prism.ingest.parsed_table import ParsedTable
from prism.schema import YesNo

PARSER_NAME = "mospi-plfs"
PARSER_VERSION = "1.0.0"
PARSER = f"{PARSER_NAME}-{PARSER_VERSION}"

RATE_COLUMNS = (
    "geography",
    "sector",
    "sex",
    "value",
    "unit",
    "measure",
    "reference_period",
    "age_group",
)
EARNINGS_COLUMNS = RATE_COLUMNS
MONTHLY_COLUMNS = RATE_COLUMNS

SECTOR_HEADER_NORM = {
    "rural": "rural",
    "urban": "urban",
    "rural + urban": "rural + urban",
    "rural+urban": "rural + urban",
}
SEX_NORM = {
    "male": "male",
    "female": "female",
    "person": "person",
}
_WITHHELD = frozenset(
    {"-", "–", "—", "..", "...", "…", "n.a.", "na", "n/a", "not available"}
)
_AGE_RE = re.compile(r"age\s*group\s*:\s*(.+?)(?:\s{2,}|\n|$)", re.IGNORECASE)
_PERIOD_QUARTER_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|"
    r"November|December)\s*[–\-]\s*"
    r"(January|February|March|April|May|June|July|August|September|October|"
    r"November|December)\s+(\d{4})",
    re.IGNORECASE,
)
_PERIOD_MONTH_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|"
    r"November|December),?\s+(\d{4})",
    re.IGNORECASE,
)
_MONTH_CELL_RE = re.compile(
    r"^(January|February|March|April|May|June|July|August|September|October|"
    r"November|December),?\s+(\d{4})$",
    re.IGNORECASE,
)
_NUMBER_RE = re.compile(r"^-?\d+(?:\.\d+)?$")
_SECTOR_CELL = frozenset({"rural", "urban", "rural + urban", "rural+urban"})


def _empty_fail(flags: str) -> ParsedTable:
    return ParsedTable(
        csv_text="",
        row_count=0,
        nulls="not parsed",
        flags=flags,
        lineage_ok=YesNo.no,
    )


def _to_csv(fieldnames: tuple[str, ...], rows: list[dict[str, str]]) -> str:
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def _ok_rows(
    fieldnames: tuple[str, ...],
    rows: list[dict[str, str]],
    value_columns: tuple[str, ...],
    flags: str,
) -> ParsedTable:
    if not rows:
        return _empty_fail(flags if flags != "none" else "empty_table")
    nulls = "; ".join(
        f"{column}: {sum(1 for row in rows if not row.get(column))}"
        for column in value_columns
    )
    return ParsedTable(
        csv_text=_to_csv(fieldnames, rows),
        row_count=len(rows),
        nulls=nulls or "none",
        flags=flags,
        lineage_ok=YesNo.yes,
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in _WITHHELD:
        return ""
    return text


def _normalize_sector(label: str) -> str | None:
    key = re.sub(r"\s+", " ", label.strip().lower())
    return SECTOR_HEADER_NORM.get(key)


def _normalize_sex(label: str) -> str | None:
    return SEX_NORM.get(label.strip().lower())


def _grid_from_bytes(data: bytes) -> list[list[object | None]]:
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows: list[list[object | None]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
        return rows
    finally:
        workbook.close()


def _title_measure_rate(title: str) -> str | None:
    upper = title.upper()
    found: list[str] = []
    for token in ("LFPR", "WPR", "UR"):
        if re.search(rf"\b{token}\b", upper) or (
            token == "UR" and "UNEMPLOYMENT RATE" in upper
        ):
            found.append(token)
        elif token == "LFPR" and "LABOUR FORCE PARTICIPATION" in upper:
            found.append("LFPR")
        elif token == "WPR" and "WORKER POPULATION" in upper:
            found.append("WPR")
    # Prefer explicit acronyms; collapse duplicates
    ordered: list[str] = []
    for item in found:
        if item not in ordered:
            ordered.append(item)
    if len(ordered) == 1:
        return ordered[0]
    return None


def _title_measure_earnings(title: str) -> str | None:
    lower = title.lower()
    hits: list[str] = []
    if "regular wage" in lower or "salaried" in lower:
        hits.append("regular-wage-salary")
    if "casual labour" in lower or "casual labor" in lower:
        hits.append("casual-labour")
    if "self-employment" in lower or "self employment" in lower:
        hits.append("self-employment")
    if len(hits) == 1:
        return hits[0]
    return None


def _age_group_from_title(title: str) -> str | None:
    match = _AGE_RE.search(title.replace("\xa0", " "))
    if match is None:
        return None
    return re.sub(r"\s+", " ", match.group(1)).strip(" .")


def _period_from_title(title: str, *, default_annual: str | None = None) -> str | None:
    cleaned = title.replace("\xa0", " ")
    quarter = _PERIOD_QUARTER_RE.search(cleaned)
    if quarter is not None:
        start, end, year = quarter.groups()
        return f"{start}-{end} {year}"
    months = list(_PERIOD_MONTH_RE.finditer(cleaned))
    if len(months) == 1:
        month, year = months[0].groups()
        return f"{month} {year}"
    if default_annual is not None and (
        "usual status" in cleaned.lower()
        or "current weekly status" in cleaned.lower()
        or "earnings" in cleaned.lower()
        or "wage/salary" in cleaned.lower()
        or "self-employment" in cleaned.lower()
        or "casual labour" in cleaned.lower()
    ):
        # Annual Appendix tables omit the calendar year in the sheet title;
        # citation card locks January–December 2025.
        return default_annual
    return None


def _is_title_row(cell: str) -> bool:
    stripped = cell.lstrip("`").strip()
    return stripped.lower().startswith("table")


def _is_geo_header(cell: str) -> bool:
    key = cell.strip().lower().replace("\\", "/")
    return key in {"state", "state/ut", "state / ut", "state\\ut"}


def _find_sector_sex_layout(
    rows: list[list[object | None]], header_idx: int
) -> tuple[list[tuple[int, str, str]], str] | tuple[None, str]:
    if header_idx + 1 >= len(rows):
        return None, "missing_sex_header_row"
    header = rows[header_idx]
    sex_row = rows[header_idx + 1]
    sector_by_col: dict[int, str] = {}
    current: str | None = None
    max_col = max(len(header), len(sex_row))
    for col in range(1, max_col):
        label = _cell_text(header[col] if col < len(header) else None)
        if label:
            sector = _normalize_sector(label)
            if sector is None:
                return None, f"unrecognised_sector_header:{label!r}"
            current = sector
        if current is not None:
            sector_by_col[col] = current
    layout: list[tuple[int, str, str]] = []
    for col, sector in sector_by_col.items():
        sex_label = _cell_text(sex_row[col] if col < len(sex_row) else None)
        if not sex_label:
            continue
        sex = _normalize_sex(sex_label)
        if sex is None:
            return None, f"unrecognised_sex_header:{sex_label!r}"
        layout.append((col, sector, sex))
    if len(layout) != 9:
        return None, f"expected_9_value_columns_got_{len(layout)}"
    expected_sectors = ["rural"] * 3 + ["urban"] * 3 + ["rural + urban"] * 3
    expected_sex = ["male", "female", "person"] * 3
    for index, (_, sector, sex) in enumerate(layout):
        if sector != expected_sectors[index] or sex != expected_sex[index]:
            return None, "sector_sex_layout_mismatch"
    return layout, "none"


def _unpivot_plfs_xlsx(
    data: bytes,
    *,
    measure_from_title,
    unit: str,
    default_period: str | None,
) -> ParsedTable:
    try:
        grid = _grid_from_bytes(data)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types
        return _empty_fail(f"xlsx_read_failed:{exc}")
    if not grid:
        return _empty_fail("empty_workbook")

    rows_out: list[dict[str, str]] = []
    flags_parts: list[str] = []
    row_i = 0
    while row_i < len(grid):
        first = _cell_text(grid[row_i][0] if grid[row_i] else None)
        if not _is_title_row(first):
            row_i += 1
            continue
        title = first
        measure = measure_from_title(title)
        if measure is None:
            return _empty_fail(f"ambiguous_or_missing_measure_in_title:{title[:80]!r}")
        age_group = _age_group_from_title(title)
        if age_group is None:
            # Earnings Table 38–40 have no age-group line — leave blank, not invent.
            age_group = ""
            if measure in {"LFPR", "WPR", "UR"}:
                return _empty_fail(f"missing_age_group_in_title:{title[:80]!r}")
        period = _period_from_title(title, default_annual=default_period)
        if period is None:
            return _empty_fail(f"missing_reference_period_in_title:{title[:80]!r}")

        header_idx = None
        for probe in range(row_i + 1, min(row_i + 6, len(grid))):
            geo_cell = _cell_text(grid[probe][0] if grid[probe] else None)
            if _is_geo_header(geo_cell):
                header_idx = probe
                break
        if header_idx is None:
            return _empty_fail("missing_state_header_after_title")

        layout, layout_flag = _find_sector_sex_layout(grid, header_idx)
        if layout is None:
            return _empty_fail(layout_flag)

        data_start = header_idx + 2
        # Skip blank rows and column-number rows like (1) (2) …
        while data_start < len(grid):
            probe_geo = _cell_text(grid[data_start][0] if grid[data_start] else None)
            if probe_geo == "":
                data_start += 1
                continue
            if probe_geo.startswith("(") and probe_geo.endswith(")"):
                data_start += 1
                continue
            break

        block_rows = 0
        for data_i in range(data_start, len(grid)):
            geo = _cell_text(grid[data_i][0] if grid[data_i] else None)
            if not geo:
                break
            if _is_title_row(geo) or geo.lower().startswith("note") or geo.lower().startswith(
                "sample"
            ):
                break
            if geo.startswith("(") and geo.endswith(")"):
                continue
            for col, sector, sex in layout:
                raw = grid[data_i][col] if col < len(grid[data_i]) else None
                value = _cell_text(raw)
                rows_out.append(
                    {
                        "geography": geo,
                        "sector": sector,
                        "sex": sex,
                        "value": value,
                        "unit": unit,
                        "measure": measure,
                        "reference_period": period,
                        "age_group": age_group,
                    }
                )
            block_rows += 1
        if block_rows == 0:
            return _empty_fail("title_block_had_no_geography_rows")
        flags_parts.append(f"{measure}:{age_group or 'no-age'}:{block_rows}geos")
        row_i = data_start + block_rows

    if not rows_out:
        return _empty_fail("no_table_blocks_found")
    return _ok_rows(RATE_COLUMNS, rows_out, ("value",), "; ".join(flags_parts))


def parse_plfs_rate_xlsx(data: bytes) -> ParsedTable:
    return _unpivot_plfs_xlsx(
        data,
        measure_from_title=_title_measure_rate,
        unit="percent",
        default_period="January-December 2025",
    )


def parse_plfs_earnings_xlsx(data: bytes) -> ParsedTable:
    return _unpivot_plfs_xlsx(
        data,
        measure_from_title=_title_measure_earnings,
        unit="rupees",
        default_period="January-December 2025",
    )


def _pdf_nonempty_cells(row: list[Any]) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for index, cell in enumerate(row):
        if cell is None:
            continue
        text = str(cell).strip()
        if text:
            out.append((index, text))
    return out


def _parse_statement_table(table: list[list[Any]], measure: str) -> list[dict[str, str]] | str:
    rows_out: list[dict[str, str]] = []
    sector: str | None = None
    age_group: str | None = None
    for row in table:
        cells = _pdf_nonempty_cells(row)
        if not cells:
            continue
        texts = [text for _, text in cells]
        joined = " ".join(texts).lower()
        if joined.startswith("statement"):
            continue
        if "survey month" in joined and "male" in joined:
            continue

        age_hit = next((t for t in texts if t.lower().startswith("age group:")), None)
        if age_hit is not None:
            age_group = re.sub(
                r"\s+", " ", age_hit.split(":", 1)[1]
            ).strip()
            # sector may share the row (e.g. urban + age group)
            for text in texts:
                norm = _normalize_sector(text)
                if norm is not None:
                    sector = norm
            continue

        sector_hits = [
            _normalize_sector(text)
            for text in texts
            if _normalize_sector(text) is not None
        ]
        month_cells = [text for text in texts if _MONTH_CELL_RE.match(text)]
        number_cells = [text for text in texts if _NUMBER_RE.match(text.replace(",", ""))]

        if sector_hits:
            if len(sector_hits) != 1:
                return f"ambiguous_sector_row:{texts!r}"
            sector = sector_hits[0]
            if not month_cells and not number_cells:
                continue

        if not month_cells:
            continue
        if len(month_cells) != 1:
            return f"ambiguous_month_row:{texts!r}"
        if len(number_cells) != 3:
            return f"expected_3_values_got_{len(number_cells)}:{texts!r}"
        if sector is None or age_group is None:
            return "data_row_before_sector_or_age_group"

        month_match = _MONTH_CELL_RE.match(month_cells[0])
        if month_match is None:
            return f"unparseable_month:{month_cells[0]!r}"
        month, year = month_match.groups()
        period = f"{month} {year}"
        for sex, value in zip(("male", "female", "person"), number_cells, strict=True):
            rows_out.append(
                {
                    "geography": "All India",
                    "sector": sector,
                    "sex": sex,
                    "value": value.replace(",", ""),
                    "unit": "percent",
                    "measure": measure,
                    "reference_period": period,
                    "age_group": age_group,
                }
            )
    if not rows_out:
        return f"statement_{measure}_empty"
    return rows_out


def parse_plfs_monthly_bulletin_pdf(data: bytes) -> ParsedTable:
    try:
        with pdfplumber.open(BytesIO(data)) as pdf:
            tables: list[list[list[Any]]] = []
            for page in pdf.pages:
                extracted = page.extract_tables() or []
                tables.extend(extracted)
    except (PdfminerException, OSError) as exc:
        return _empty_fail(f"pdf_open_failed:{exc}")

    wanted = {
        1: "LFPR",
        2: "WPR",
        3: "UR",
    }
    found: dict[int, list[list[Any]]] = {}
    for table in tables:
        if not table:
            continue
        first_cells = _pdf_nonempty_cells(table[0])
        if not first_cells:
            continue
        title = first_cells[0][1]
        match = re.match(r"Statement\s+(\d)\s*:", title, re.IGNORECASE)
        if match is None:
            continue
        number = int(match.group(1))
        if number not in wanted:
            continue
        if number in found:
            return _empty_fail(f"duplicate_statement_{number}")
        found[number] = table

    missing = [n for n in wanted if n not in found]
    if missing:
        return _empty_fail(f"missing_statements:{missing}")

    rows_out: list[dict[str, str]] = []
    for number, measure in wanted.items():
        parsed = _parse_statement_table(found[number], measure)
        if isinstance(parsed, str):
            return _empty_fail(parsed)
        rows_out.extend(parsed)

    return _ok_rows(
        MONTHLY_COLUMNS,
        rows_out,
        ("value",),
        "statements_1_2_3_all_india_cws",
    )


register_parser(
    "plfs-rate-xlsx",
    ParserSpec(parse=parse_plfs_rate_xlsx, lineage_name=PARSER, kind="xlsx"),
)
register_parser(
    "plfs-earnings-xlsx",
    ParserSpec(parse=parse_plfs_earnings_xlsx, lineage_name=PARSER, kind="xlsx"),
)
register_parser(
    "plfs-monthly-bulletin-pdf",
    ParserSpec(
        parse=parse_plfs_monthly_bulletin_pdf, lineage_name=PARSER, kind="pdf"
    ),
)

__all__ = [
    "PARSER",
    "parse_plfs_earnings_xlsx",
    "parse_plfs_monthly_bulletin_pdf",
    "parse_plfs_rate_xlsx",
]
